from __future__ import annotations

import csv
import hashlib
import json
import math
import re
from pathlib import Path
from typing import Any, Iterable

from docx import Document
from pypdf import PdfReader


MARKER_RE = re.compile(r"\[[A-Z][A-Z0-9_]*(?:REQUIRED|PENDING)[A-Z0-9_]*\]")


def stable(prefix: str, *parts: object, n: int = 16) -> str:
    payload = "\x1f".join(str(part) for part in parts)
    return f"{prefix}-{hashlib.sha256(payload.encode('utf-8')).hexdigest()[:n].upper()}"


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def read_csv(path: Path) -> list[dict[str, str]]:
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle))


def write_csv(path: Path, records: list[dict[str, Any]], fields: list[str] | None = None) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    fields = fields or (list(records[0]) if records else ["status"])
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(records)


def parse_scalar(value: str) -> Any:
    value = value.strip()
    if len(value) >= 2 and value[0] == value[-1] and value[0] in {'"', "'"}:
        return value[1:-1]
    if value.lower() in {"true", "false"}:
        return value.lower() == "true"
    if re.fullmatch(r"-?\d+", value):
        return int(value)
    if re.fullmatch(r"-?\d+\.\d+", value):
        return float(value)
    return value


def parse_brief(path: Path) -> dict[str, Any]:
    """Parse the deliberately small, flat brief YAML contract without PyYAML."""
    lines = path.read_text(encoding="utf-8-sig").splitlines()
    result: dict[str, Any] = {}
    i = 0
    while i < len(lines):
        raw = lines[i]
        if not raw.strip() or raw.lstrip().startswith("#") or raw.startswith(" "):
            i += 1
            continue
        key, sep, value = raw.partition(":")
        if not sep:
            i += 1
            continue
        key = key.strip()
        value = value.strip()
        if value in {">", "|"}:
            parts: list[str] = []
            i += 1
            while i < len(lines) and (lines[i].startswith("  ") or not lines[i].strip()):
                if lines[i].strip():
                    parts.append(lines[i].strip())
                i += 1
            result[key] = " ".join(parts)
            continue
        if value == "":
            items: list[str] = []
            i += 1
            while i < len(lines) and (lines[i].startswith("  ") or not lines[i].strip()):
                stripped = lines[i].strip()
                if stripped.startswith("- "):
                    items.append(stripped[2:].strip())
                i += 1
            result[key] = items
            continue
        result[key] = parse_scalar(value)
        i += 1
    return result


def extract_text(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix in {".md", ".txt", ".ris", ".csv", ".tsv", ".yaml", ".yml"}:
        return path.read_text(encoding="utf-8-sig", errors="replace")
    if suffix == ".docx":
        doc = Document(path)
        content = [paragraph.text for paragraph in doc.paragraphs]
        for table in doc.tables:
            content.extend(" | ".join(cell.text for cell in row.cells) for row in table.rows)
        return "\n".join(content)
    if suffix == ".pdf":
        return "\n".join(page.extract_text() or "" for page in PdfReader(path).pages)
    return ""


ROLE_SIGNATURES: list[tuple[str, set[str]]] = [
    ("balance", {"covariate", "absolute_smd_unweighted", "absolute_smd_weighted"}),
    ("effects", {"claim_id", "estimand", "contrast", "estimate", "ci_low", "ci_high"}),
    ("sensitivity", {"analysis", "risk_difference", "ci_low", "ci_high"}),
    ("weight_distribution", {"stabilized_weight"}),
    ("cell_metadata", {"cell_id", "group", "cell_type", "UMAP_1", "UMAP_2"}),
    ("composition", {"group", "cell_type", "cell_count", "percentage"}),
    ("deg", {"gene", "cell_type", "FDR"}),
    ("pathway", {"pathway", "cell_type", "FDR", "evidence_level"}),
    ("communication", {"sender", "receiver", "ligand_receptor", "communication_probability"}),
    ("qc", {"stage", "remaining_cells", "removed_at_stage"}),
    ("study_level", {"study_id", "study_label", "sample_size", "odds_ratio", "ci_low", "ci_high"}),
    ("pooled", {"claim_id", "estimand", "estimate", "I2_percent", "study_count"}),
    ("risk_of_bias", {"study_id", "study_label", "domain", "judgment"}),
    ("subgroup", {"subgroup", "pooled_or", "ci_low", "ci_high", "interaction_p"}),
    ("grade", {"outcome", "studies", "participants", "certainty"}),
    ("eligibility", {"category", "criterion", "operational_status"}),
    ("schedule", {"visit", "assessment", "window", "status"}),
    ("sample_size", {"parameter", "value", "unit"}),
    ("variables", {"role", "variable", "timepoint", "type"}),
    ("risk_register", {"risk", "domain", "probability", "impact", "mitigation"}),
    ("milestones", {"milestone", "start_date", "end_date", "status"}),
]


def classify_dataset(columns: Iterable[str]) -> str:
    field_set = set(columns)
    if field_set == {"stabilized_weight"}:
        return "weight_distribution"
    for role, signature in ROLE_SIGNATURES:
        if role == "weight_distribution":
            continue
        if signature <= field_set:
            if role == "deg" and not any(field.lower().startswith("log2_fold_change") for field in field_set):
                continue
            if role == "pathway" and not any("enrichment_score" in field.lower() for field in field_set):
                continue
            return role
    return "unclassified"


def discover_project(project: Path) -> tuple[dict[str, Any], list[dict[str, Any]], dict[str, dict[str, Any]], dict[str, str]]:
    brief = parse_brief(project / "brief" / "presentation_brief.yaml")
    sources: list[dict[str, Any]] = []
    datasets: dict[str, dict[str, Any]] = {}
    texts: dict[str, str] = {}
    candidates = [
        project / "README_PROJECT.md",
        project / "brief" / "presentation_brief.yaml",
        *sorted((project / "input").rglob("*")),
    ]
    for source in candidates:
        if not source.is_file() or "expected_ground_truth" in source.parts:
            continue
        relative = source.relative_to(project).as_posix()
        source_id = stable("SRC", project.name, relative, n=12)
        text = extract_text(source)
        if text:
            texts[relative] = text
        role = "document"
        parsed_rows: list[dict[str, str]] = []
        columns: list[str] = []
        if source.suffix.lower() == ".csv":
            parsed_rows = read_csv(source)
            columns = list(parsed_rows[0]) if parsed_rows else []
            role = classify_dataset(columns)
            if role != "unclassified":
                if role in datasets:
                    raise ValueError(f"Duplicate canonical dataset role {role}: {source}")
                datasets[role] = {
                    "role": role,
                    "source_id": source_id,
                    "source_file": relative,
                    "absolute_path": str(source.resolve()),
                    "columns": columns,
                    "rows": parsed_rows,
                }
        elif source.suffix.lower() in {".png", ".jpg", ".jpeg", ".svg"}:
            role = "figure"
        elif source.suffix.lower() in {".ris"}:
            role = "reference"
        elif source.suffix.lower() in {".yaml", ".yml"}:
            role = "brief"
        sources.append({
            "source_id": source_id,
            "source_file": relative,
            "source_name": source.name,
            "source_type": source.suffix.lower(),
            "source_role": role,
            "sha256": sha256(source),
            "size_bytes": source.stat().st_size,
            "parsed": bool(text or source.suffix.lower() in {".xlsx", ".png", ".jpg", ".jpeg", ".svg"}),
            "ground_truth_excluded": True,
        })
    return brief, sources, datasets, texts


def infer_kind(datasets: dict[str, dict[str, Any]]) -> str:
    roles = set(datasets)
    if {"balance", "effects", "sensitivity"} <= roles:
        return "target_trial"
    if {"composition", "deg", "pathway", "communication", "qc"} <= roles:
        return "single_cell"
    if {"study_level", "pooled", "risk_of_bias", "subgroup", "grade"} <= roles:
        return "meta_analysis"
    if {"eligibility", "schedule", "sample_size", "variables", "risk_register", "milestones"} <= roles:
        return "protocol"
    raise ValueError(f"Unable to infer project type from dataset roles: {sorted(roles)}")


def source_binding(
    dataset: dict[str, Any],
    fields_used: list[str],
    source_location: str = "all rows",
    row_filter: str = "all rows",
    aggregation: str = "none",
    canonical_status: str = "canonical",
) -> dict[str, Any]:
    missing = [field for field in fields_used if field not in dataset["columns"]]
    if missing:
        raise ValueError(f"Unbound fields for {dataset['role']}: {missing}")
    return {
        "source_id": dataset["source_id"],
        "source_file": dataset["source_file"],
        "source_location": source_location,
        "fields_used": fields_used,
        "row_filter": row_filter,
        "aggregation": aggregation,
        "canonical_status": canonical_status,
        "short_label_zh": short_source_label(dataset["role"]),
    }


def document_binding(sources: list[dict[str, Any]], keyword: str, label: str) -> dict[str, Any]:
    candidates = [source for source in sources if source["source_role"] == "document"]
    match = next((source for source in candidates if keyword.lower() in source["source_name"].lower()), None)
    match = match or (candidates[0] if candidates else None)
    if match is None:
        raise ValueError("No document source available for narrative binding")
    return {
        "source_id": match["source_id"],
        "source_file": match["source_file"],
        "source_location": label,
        "fields_used": ["document_text"],
        "row_filter": "relevant structured section",
        "aggregation": "none",
        "canonical_status": "supporting",
        "short_label_zh": "研究方案" if "protocol" in keyword.lower() else "补充方法",
    }


def asset_binding(sources: list[dict[str, Any]], keyword: str, label: str) -> dict[str, Any]:
    candidates = [source for source in sources if source["source_role"] == "figure"]
    match = next((source for source in candidates if keyword.lower() in source["source_name"].lower()), None)
    if match is None:
        raise ValueError(f"No registered figure contains keyword: {keyword}")
    return {
        "source_id": match["source_id"],
        "source_file": match["source_file"],
        "source_location": label,
        "fields_used": ["registered_figure"],
        "row_filter": "entire registered figure",
        "aggregation": "none",
        "canonical_status": "supporting",
        "short_label_zh": "登记图形",
    }


def short_source_label(role: str) -> str:
    labels = {
        "balance": "平衡诊断",
        "effects": "冻结效应估计",
        "sensitivity": "敏感性分析",
        "weight_distribution": "权重分布",
        "cell_metadata": "单细胞元数据",
        "composition": "细胞组成表",
        "deg": "Pseudobulk差异分析",
        "pathway": "通路富集表",
        "communication": "细胞通讯结果",
        "qc": "QC汇总表",
        "study_level": "研究级数据",
        "pooled": "合并效应",
        "risk_of_bias": "偏倚风险表",
        "subgroup": "亚组Meta分析",
        "grade": "GRADE汇总",
        "eligibility": "入排标准",
        "schedule": "评估时间表",
        "sample_size": "样本量假设",
        "variables": "变量字典",
        "risk_register": "风险登记表",
        "milestones": "项目里程碑",
    }
    return labels.get(role, "输入材料")


def detect_conflicts(project_name: str, sources: list[dict[str, Any]], datasets: dict[str, dict[str, Any]], texts: dict[str, str]) -> list[dict[str, Any]]:
    source_by_file = {source["source_file"]: source for source in sources}
    found: dict[tuple[str, str], dict[str, Any]] = {}
    for relative, text in texts.items():
        for line_no, line in enumerate(text.splitlines(), 1):
            match = re.search(r"(?:Use|Report)\s+(.+?),\s+not\s+(.+?)(?:\s+value)?[.!]?$", line.strip().lstrip("- ").strip(), re.I)
            if not match or not re.search(r"\d", match.group(1)) or not re.search(r"\d", match.group(2)):
                continue
            canonical, conflicting = match.group(1).strip(), match.group(2).strip()
            found[(conflicting, canonical)] = {
                "conflict_id": stable("CONFLICT", project_name, relative, line_no, n=12),
                "source_id": source_by_file[relative]["source_id"],
                "source_file": relative,
                "source_location": f"line {line_no}",
                "conflicting_value": conflicting,
                "canonical_value": canonical,
                "resolution": "保留冲突记录，并使用已声明的canonical值。",
                "manual_review_required": True,
            }
    composition = datasets.get("composition")
    if composition:
        # Source prose may localize biological entity names while the canonical
        # table keeps English labels.  Match by group plus an explicit entity
        # alias; never compare every percentage in a paragraph to every row.
        # This compact vocabulary is domain terminology, not a project/title
        # special case, and can be extended without changing conflict logic.
        cell_type_aliases = {
            "macrophage": {"macrophage", "macrophages", "巨噬细胞"},
            "monocyte": {"monocyte", "monocytes", "单核细胞"},
            "t cell": {"t cell", "t cells", "t细胞"},
            "b cell": {"b cell", "b cells", "b细胞"},
            "nk cell": {"nk cell", "nk cells", "nk细胞"},
            "endothelial": {"endothelial", "endothelial cell", "内皮细胞"},
            "fibroblast": {"fibroblast", "fibroblasts", "成纤维细胞"},
            "dendritic": {"dendritic", "dendritic cell", "树突状细胞"},
        }
        for canonical in composition.get("rows", []):
            group = str(canonical.get("group", "")).strip()
            cell_type = str(canonical.get("cell_type", "")).strip()
            canonical_percentage = canonical.get("percentage")
            if not group or not cell_type or canonical_percentage in (None, ""):
                continue
            aliases = cell_type_aliases.get(cell_type.casefold(), {cell_type.casefold()})
            canonical_number = float(canonical_percentage)
            for relative, text in texts.items():
                source = source_by_file.get(relative)
                if not source:
                    continue
                for line_no, line in enumerate(text.splitlines(), 1):
                    normalized = line.casefold().replace(" ", "")
                    if group.casefold().replace(" ", "") not in normalized:
                        continue
                    if not any(alias.casefold().replace(" ", "") in normalized for alias in aliases):
                        continue
                    for match in re.finditer(r"(\d+(?:\.\d+)?)\s*%", line):
                        observed_number = float(match.group(1))
                        if math.isclose(observed_number, canonical_number, rel_tol=0, abs_tol=0.005):
                            continue
                        conflicting = f"{match.group(1)}%"
                        correct = f"{canonical_percentage}%"
                        found[(conflicting, correct)] = {
                            "conflict_id": stable("CONFLICT", project_name, relative, line_no, n=12),
                            "source_id": source["source_id"],
                            "source_file": relative,
                            "source_location": f"line {line_no}",
                            "conflicting_value": conflicting,
                            "canonical_value": correct,
                            "resolution": "使用canonical细胞组成表中的分组比例。",
                            "manual_review_required": True,
                        }
    return list(found.values())


def detect_unresolved(sources: list[dict[str, Any]], texts: dict[str, str]) -> list[dict[str, Any]]:
    source_by_file = {source["source_file"]: source for source in sources}
    found: dict[str, dict[str, Any]] = {}
    for relative, text in texts.items():
        for marker in MARKER_RE.findall(text):
            found[marker] = {
                "item_id": stable("UNRES", marker, n=12),
                "marker": marker,
                "status": "INFORMATION_REQUIRED",
                "source_id": source_by_file[relative]["source_id"],
                "source_file": relative,
                "source_location": "marker occurrence",
                "required_action": "等待用户确认，不得虚构补全。",
                "manual_review_required": True,
            }
    return list(found.values())


VISUAL_REQUIRED = {
    "forest_plot": ["x_field", "ci_low_field", "ci_high_field", "label_field", "reference_value"],
    "love_plot": ["x_field", "y_field", "label_field", "reference_value"],
    "grouped_composition": ["y_field", "group_field", "label_field"],
    "diverging_enrichment": ["x_field", "facet_field", "label_field", "reference_value"],
    "communication_network": ["edge_source_field", "edge_target_field", "label_field"],
    "risk_of_bias_matrix": ["group_field", "facet_field", "color_field"],
    "risk_matrix": ["x_field", "y_field", "label_field"],
    "timeline_gantt": ["time_field", "label_field"],
}


def validate_visual_spec(spec: dict[str, Any]) -> list[str]:
    required = {
        "visual_id", "visual_type", "scientific_role", "data_contract_id",
        "source_bindings", "row_filter", "aggregation", "entity_scope",
        "expected_row_count", "rendered_row_count", "expected_row_keys", "rendered_row_keys",
        "top_k_policy", "omitted_rows",
        "editability_requirement", "design_variant",
    }
    errors = [f"missing:{field}" for field in sorted(required - set(spec))]
    if not spec.get("source_bindings"):
        errors.append("source_bindings_empty")
    rendered = spec.get("rendered_row_count")
    if rendered is not None and spec.get("expected_row_count") != rendered and not spec.get("omitted_rows"):
        errors.append("silent_truncation")
    expected_keys = spec.get("expected_row_keys", [])
    rendered_keys = spec.get("rendered_row_keys", [])
    if len(expected_keys) != spec.get("expected_row_count") or len(set(expected_keys)) != len(expected_keys):
        errors.append("expected_row_identity_invalid")
    if rendered is not None and set(expected_keys) != set(rendered_keys) and not spec.get("omitted_rows"):
        errors.append("rendered_row_identity_mismatch")
    for field in VISUAL_REQUIRED.get(spec.get("visual_type", ""), []):
        if spec.get(field) in {None, ""}:
            errors.append(f"chart_field_missing:{field}")
    if spec.get("visual_type") == "communication_network" and spec.get("sequential") is not False:
        errors.append("network_marked_sequential")
    return errors


def validate_slide_spec(spec: dict[str, Any]) -> list[str]:
    required = {
        "slide_id", "slide_role", "narrative_purpose", "title", "language",
        "key_message", "claim_bindings", "source_bindings", "conflict_bindings",
        "unresolved_bindings", "visual_specs", "layout_family", "design_profile",
        "appendix_status", "review_flags",
    }
    errors = [f"missing:{field}" for field in sorted(required - set(spec))]
    if spec.get("language") not in {"zh-CN", "en-US"}:
        errors.append("unsupported_language")
    for visual in spec.get("visual_specs", []):
        errors.extend(f"{visual.get('visual_id', 'unknown')}:{error}" for error in validate_visual_spec(visual))
    return errors


def validate_visual_semantics(spec: dict[str, Any]) -> list[str]:
    """Validate scientific meaning encoded in a resolved VisualSpec payload."""
    errors: list[str] = []
    visual_type = spec.get("visual_type", "")
    payload = spec.get("payload", {})
    rows = payload.get("rows", [])
    if visual_type in {"forest_plot", "subgroup_forest"}:
        if spec.get("reference_value") is None:
            errors.append("forest_reference_missing")
        if not payload.get("axis_label"):
            errors.append("forest_axis_missing")
        if payload.get("scale") == "log":
            for row in rows:
                values = [row.get(spec.get("x_field")), row.get(spec.get("ci_low_field")), row.get(spec.get("ci_high_field"))]
                if any(float(value) <= 0 for value in values if value not in {None, ""}):
                    errors.append("log_scale_non_positive")
                    break
        for row in rows:
            try:
                estimate = float(row[spec["x_field"]])
                low = float(row[spec["ci_low_field"]])
                high = float(row[spec["ci_high_field"]])
                if not low < estimate < high:
                    errors.append("invalid_confidence_interval")
                    break
            except (KeyError, TypeError, ValueError):
                errors.append("forest_numeric_field_missing")
                break
    if visual_type == "pooled_evidence_panel":
        prediction = payload.get("prediction", {})
        if "prediction" not in prediction.get("estimand", "").lower():
            errors.append("prediction_interval_not_explicit")
    if visual_type == "love_plot":
        required = {"absolute_smd_unweighted", "absolute_smd_weighted"}
        if not rows or not all(required <= set(row) for row in rows):
            errors.append("love_before_after_missing")
    if visual_type == "grouped_composition":
        groups = {row.get(spec.get("group_field")) for row in rows}
        if len(groups - {None, ""}) < 2:
            errors.append("composition_group_comparison_missing")
    if visual_type == "diverging_enrichment":
        values = [float(row[spec["x_field"]]) for row in rows]
        if not any(value < 0 for value in values) or not any(value > 0 for value in values):
            errors.append("diverging_sign_not_represented")
        facets = {row.get(spec.get("facet_field")) for row in rows}
        if len(facets - {None, ""}) < 2:
            errors.append("enrichment_facets_missing")
    if visual_type == "communication_network":
        if spec.get("sequential") is not False:
            errors.append("independent_network_marked_sequence")
        edge_pairs = {(row.get(spec.get("edge_source_field")), row.get(spec.get("edge_target_field"))) for row in rows}
        if len(edge_pairs) != len(rows):
            errors.append("network_edge_identity_lost")
    if visual_type == "risk_of_bias_matrix":
        studies = {row.get(spec.get("group_field")) for row in rows}
        domains = {row.get(spec.get("facet_field")) for row in rows}
        pairs = {(row.get(spec.get("group_field")), row.get(spec.get("facet_field"))) for row in rows}
        if len(pairs) != len(studies) * len(domains):
            errors.append("rob_matrix_incomplete")
    if visual_type == "risk_matrix":
        if not rows or any(row.get(spec.get("x_field")) in {None, ""} or row.get(spec.get("y_field")) in {None, ""} for row in rows):
            errors.append("risk_matrix_axes_missing")
    if visual_type == "dag":
        nodes = {node.get("id"): node for node in payload.get("nodes", [])}
        edges = {(edge.get("source"), edge.get("target")) for edge in payload.get("edges", [])}
        exposures = [key for key, node in nodes.items() if node.get("role") == "exposure"]
        outcomes = [key for key, node in nodes.items() if node.get("role") == "outcome"]
        confounders = [key for key, node in nodes.items() if node.get("role") == "confounder"]
        if not exposures or not outcomes or not confounders:
            errors.append("dag_semantic_roles_missing")
        else:
            exposure, outcome = exposures[0], outcomes[0]
            if (exposure, outcome) not in edges:
                errors.append("dag_exposure_outcome_edge_missing")
            for confounder in confounders:
                if (confounder, exposure) not in edges or (confounder, outcome) not in edges:
                    errors.append("dag_confounder_paths_missing")
                    break
    if visual_type == "timeline_gantt":
        starts = [row.get("start_date") for row in rows]
        ends = [row.get("end_date") for row in rows]
        if not rows or len(set(starts + ends)) < 3:
            errors.append("gantt_dates_not_variable")
    return errors


def validate_slide_semantics(spec: dict[str, Any]) -> list[str]:
    errors: list[str] = []
    title = spec.get("title", "")
    for visual in spec.get("visual_specs", []):
        errors.extend(f"{visual.get('visual_id', 'unknown')}:{error}" for error in validate_visual_semantics(visual))
        if visual.get("facet_field") and visual.get("payload", {}).get("rows"):
            facets = {row.get(visual["facet_field"]) for row in visual["payload"]["rows"]}
            specific = [facet for facet in facets if facet and str(facet).lower() in title.lower()]
            if specific and len(facets) > 1 and "分层" not in title and "不同" not in title:
                errors.append(f"title_entity_scope_mismatch:{specific[0]}")
    return errors


def visible_cjk_ratio(texts: Iterable[str]) -> float:
    combined = "".join(str(text) for text in texts)
    cjk = len(re.findall(r"[\u3400-\u9fff]", combined))
    alphabetic = len(re.findall(r"[A-Za-z\u3400-\u9fff]", combined))
    return cjk / alphabetic if alphabetic else 1.0


def merge_bindings(items: Iterable[dict[str, Any]]) -> list[dict[str, Any]]:
    unique: dict[tuple[str, tuple[str, ...], str, str], dict[str, Any]] = {}
    for item in items:
        key = (
            item["source_id"],
            tuple(item.get("fields_used", [])),
            item.get("row_filter", ""),
            item.get("aggregation", ""),
        )
        unique[key] = item
    return list(unique.values())


def dump_json(path: Path, value: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2), encoding="utf-8")
