"""Isolated synthetic shadow experiment; never a public production entry point.

Run with Python -B. The only accepted content is the checked-in synthetic fixture.
All artifacts go into a new output directory. Scientific modules remain read-only.
"""
from __future__ import annotations

import argparse
from collections import Counter
import copy
import csv
import hashlib
import io
import json
import os
from pathlib import Path
import re
import secrets
import shutil
import subprocess
import sys
import time
import unittest
import zipfile
from xml.etree import ElementTree as ET

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "scripts"))
from academic_ppt.chart_data import ChartDataContract
from academic_ppt.deck_ir import build_deck_ir, canonical_json_hash, to_backend_spec, to_storyboard_rows
from academic_ppt.evidence import build_evidence
from academic_ppt.extractors import extract_all
from academic_ppt.inventory import inventory_sources, update_manifest, verify_input_hashes
from academic_ppt.layout_contract import LayoutContract, plan_slide_geometry, validate_planned_geometry
from academic_ppt.ooxml_qa import inspect_package
from academic_ppt.qa import run_qa, write_traceability_maps
from academic_ppt.rendering import create_contact_sheet
from academic_ppt.visual_density import write_density_reports
from academic_ppt.visual_layout_qa import inspect_powerpoint_text_layout, inspect_text_geometry
from audit_repository_content import _scan_office_package
from pptx import Presentation

ID = "VISUAL_FIDELITY_EXPERIMENT_001"
BASELINE = "4e28861724938cb0e3e90ade86e3b1012a4ff6ff"
FIXTURE = ROOT / "tests/fixtures/synthetic/visual_fidelity_001.json"
GENERATOR = ROOT / "scripts/generate_deck_pptxgen.mjs"
ALLOWED = ["scripts/generate_deck_pptxgen.mjs", "scripts/experiments/visual_fidelity_001.mjs",
           "scripts/experiments/visual_fidelity_001.py", "tests/test_visual_fidelity_experiment_001.py",
           "tests/fixtures/synthetic/visual_fidelity_001.json"]
BASELINE_MODULES = [
    "tests.test_v2_7_visual_brief_contract", "tests.test_v2_ir_routing", "tests.test_v2_style_ooxml",
    "tests.test_v2_visual_modules", "tests.test_v2_visual_layout_qa",
    *["tests.ppt_regression.test_" + n for n in ("chart_metadata", "cross_renderer", "density_budget",
        "footer_intrusion", "layout_repetition", "safe_zones", "text_bounds", "visual_router")],
    "tests.art_direction.test_v2_5_art_direction_contracts", "tests.presentation_design.test_v2_4_presentation_contracts",
    "tests.test_v2_7_project_interface", "tests.test_v2_7_1_repository_privacy",
]
PRESERVATION_MODULES = ["tests.test_v2_incremental_template", "tests.fast_enhance.test_keep_semantic_preservation"]
NS = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}


def write_json(path: Path, value) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("x", encoding="utf-8") as handle:
        json.dump(value, handle, ensure_ascii=False, indent=2)
        handle.write("\n")


def sha(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def gate(condition, message: str) -> None:
    if not condition:
        raise RuntimeError(message)


def prepare(root: Path) -> dict:
    """Register, extract, and freeze one synthetic source package for BOTH arms."""
    fixture = json.loads(FIXTURE.read_text(encoding="utf-8"))
    gate(fixture["synthetic_only"] is True and fixture["fixture_id"] == ID, "Not the authorized synthetic fixture")
    source = root / "input"
    staging = root / "staging"
    source.mkdir(parents=True, exist_ok=False)
    staging.mkdir()
    shutil.copyfile(FIXTURE, source / "fixture.json")
    (source / "figure.svg").write_text(fixture["source_figure_svg"], encoding="utf-8")
    with (source / "values.csv").open("x", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        writer.writerow(["category", "rate_percent"])
        writer.writerows(zip(fixture["chart"]["categories"], fixture["chart"]["values"], strict=True))
    lines = ["Fully synthetic software fixture. " + fixture["boundary"]]
    for slide in fixture["slides"]:
        lines.append(f"CLAIM|{slide['message']}|fixture {slide['key']}|synthetic software fixture|high|{slide['message']}|Do not infer clinical effects")
        lines.extend([slide["title"], *slide.get("takeaways", [])])
        if "diagram_spec" in slide:
            lines.append(json.dumps(slide["diagram_spec"], sort_keys=True))
    (source / "statements.md").write_text("\n".join(lines) + "\n", encoding="utf-8")
    manifest = inventory_sources(source, {".md", ".svg", ".csv"}, staging / "source_manifest.csv")
    extracted = extract_all(source, staging, manifest)
    update_manifest(staging / "source_manifest.csv", manifest)
    claims, unresolved = build_evidence(extracted, manifest, staging, fixture["brief"])
    gate(not unresolved, "Scientific source preparation failed: " + repr(unresolved))
    by_file = {r["file_name"]: r["source_id"] for r in manifest}
    contract = LayoutContract.from_files(ROOT / "config/layout_contract.yaml", ROOT / "config/safe_zones.yaml")
    chart = ChartDataContract(chart_id="CHT-SYNTHETIC-001", content_type="event_rates", source_id=by_file["values.csv"],
                              source_path="values.csv", **fixture["chart"]).to_dict()
    planned = []
    for index, f in enumerate(fixture["slides"]):
        s = {"logical_slide_key": f["key"], "slide_id": f["key"], "slide_role": "content", "section_id": "synthetic",
             "slide_title": f["title"], "single_key_message": f["message"], "slide_purpose": "Synthetic composition validation",
             "visual_type": f["visual_type"], "layout_family": f["visual_type"], "proposed_layout": f["visual_type"],
             "source_ids": [by_file["statements.md"]], "claim_ids": [claims[index]["claim_id"]],
             "citation_requirement": "required", "short_source_label": "Source: synthetic fixture | software validation only",
             "speaker_note_summary": fixture["boundary"] + " " + f["message"], "prohibited_overstatement": "Do not infer clinical effects",
             "manual_review_required": "yes", "confidence": "high", "editable_object_requirements": ["native_text"]}
        for key in ("takeaways", "diagram_spec", "chart_caption"):
            if key in f:
                s[key] = copy.deepcopy(f[key])
        if f["visual_type"] == "editable_bar_chart":
            s.update(chart_data=chart, editable_object_requirements=["chart", "native_text"])
            s["source_ids"].append(by_file["values.csv"])
        if f["visual_type"] == "source_figure":
            s.update(visual_asset_path="input/figure.svg", figure_ids=["FIG-0001"], editable_object_requirements=["image", "native_text"])
            s["source_ids"].append(by_file["figure.svg"])
        s["planned_geometry"] = plan_slide_geometry(s, contract)
        gate(not validate_planned_geometry(s["planned_geometry"], contract), f"Planned geometry failed: {f['key']}")
        planned.append(s)
    ir = build_deck_ir(fixture["brief"], planned)
    canonical = to_backend_spec(ir, fixture["brief"], "restrained_biomedical")
    canonical["layout_contract"] = contract.serializable()
    assignments = {s["slide_id"]: copy.deepcopy(f["assignment"]) for s, f in zip(ir["slides"], fixture["slides"], strict=True) if "assignment" in f}
    write_json(staging / "deck_ir.json", ir)
    write_json(staging / "canonical_spec.json", canonical)
    write_json(staging / "freeze.json", {"fixture_sha256": sha(FIXTURE), "canonical_hash": ir["canonical_hash"],
               "canonical_spec_sha256": sha(staging / "canonical_spec.json"), "source_hashes": {r["source_id"]: r["sha256"] for r in manifest}})
    return dict(fixture=fixture, manifest=manifest, claims=claims, ir=ir, canonical=canonical, assignments=assignments,
                input_root=source, staging=staging, contract=contract)


def arm_spec(bundle: dict, enabled: bool) -> dict:
    spec = copy.deepcopy(bundle["canonical"])
    spec["experiment"] = {"id": ID, "enabled": enabled}
    if enabled:
        spec["visual_execution_plan"] = copy.deepcopy(bundle["assignments"])
    return spec


def generate(spec: dict, arm: Path, script: Path = GENERATOR) -> Path:
    arm.mkdir(parents=True, exist_ok=False)
    write_json(arm / "build_spec.json", spec)
    pptx = arm / "deck.pptx"
    node = shutil.which("node")
    gate(node, "Node.js is required")
    env = dict(os.environ, NODE_PATH=str(ROOT / "node_modules"), PYTHONDONTWRITEBYTECODE="1")
    result = subprocess.run([node, str(script), str(arm / "build_spec.json"), str(pptx), str(arm / "preview")],
                            cwd=arm.parent, env=env, capture_output=True, text=True, encoding="utf-8", errors="replace", timeout=120)
    (arm / "generation.log").write_text(result.stdout + result.stderr, encoding="utf-8")
    gate(result.returncode == 0, "Native generation failed: " + result.stderr)
    return pptx


def normalized_tokens(texts: list[str]) -> Counter:
    """Whitespace/box-order independent token multiset; retain punctuation/numbers.

    Paired with exact frozen-paragraph coverage below; token equality alone cannot
    prove that recommendations retain their word order or meaning.
    """
    return Counter(token for text in texts for token in re.findall(r"\w+|[^\w\s]", text, re.UNICODE))


def normalized(text: str) -> str:
    return " ".join(text.split())


def workbook_content_hash(data: bytes) -> str:
    # Office creation/modified timestamps and ZIP timestamps are nondeterministic
    # packaging metadata. Compare every decompressed workbook member, removing
    # ONLY the two core timestamp values; cell values/formulas/styles remain exact.
    with zipfile.ZipFile(io.BytesIO(data)) as workbook:
        members = {}
        for name in workbook.namelist():
            raw = workbook.read(name)
            if name == "docProps/core.xml":
                xml = ET.fromstring(raw)
                for tag in ("created", "modified"):
                    for element in xml.findall("{http://purl.org/dc/terms/}" + tag):
                        element.text = "TIMESTAMP"
                raw = ET.tostring(xml)
            members[name] = hashlib.sha256(raw).hexdigest()
    return canonical_json_hash(members)


def package_snapshot(path: Path) -> dict:
    deck = Presentation(path)
    slides = []
    for slide in deck.slides:
        texts, objects = [], []
        for shape in slide.shapes:
            text = shape.text if shape.has_text_frame else ""
            if text:
                texts.append(text)
            preset = shape._element.find(".//a:prstGeom", NS)
            kind = ("chart" if shape.has_chart else "image" if shape.shape_type == 13 else
                    "text" if text else preset.get("prst") if preset is not None else "other")
            if shape._element.tag.endswith("}cxnSp"):
                kind = "line"
            objects.append({"name": shape.name, "kind": kind, "text": text,
                            "bounds": {"x": shape.left / 914400, "y": shape.top / 914400,
                                       "w": shape.width / 914400, "h": shape.height / 914400}})
        slides.append({"texts": texts, "tokens": dict(normalized_tokens(texts)), "objects": objects})
    with zipfile.ZipFile(path) as z:
        names = z.namelist()
        parts = {n: hashlib.sha256(z.read(n)).hexdigest() for n in names if
                 re.fullmatch(r"ppt/(slides/slide\d+\.xml|notesSlides/notesSlide\d+\.xml|charts/.*|embeddings/.*|media/.*)", n)}
        for n in parts:
            if n.startswith("ppt/embeddings/") and n.endswith(".xlsx"):
                parts[n] = workbook_content_hash(z.read(n))
        notes = {n: [e.text or "" for e in ET.fromstring(z.read(n)).findall(".//a:t", NS)] for n in names if re.fullmatch(r"ppt/notesSlides/notesSlide\d+\.xml", n)}
    return {"slides": slides, "parts": parts, "notes": notes}


def coverage(rectangles: list[dict]) -> float:
    """Union of axis-aligned object bounds, not ink coverage or aesthetic score."""
    xs = sorted({x for r in rectangles for x in (r["x"], r["x"] + r["w"])})
    area = 0.0
    for left, right in zip(xs, xs[1:]):
        ys = sorted((r["y"], r["y"] + r["h"]) for r in rectangles if r["x"] < right and r["x"] + r["w"] > left)
        total, end = 0.0, float("-inf")
        for lo, hi in ys:
            total += max(0, hi - max(lo, end))
            end = max(end, hi)
        area += (right - left) * total
    return area


def metrics(snapshot: dict, bundle: dict, execution: dict | None) -> dict:
    families, rows = [], []
    for index, (slide, canonical) in enumerate(zip(snapshot["slides"], bundle["ir"]["slides"], strict=True)):
        exp = next((r for r in (execution or {}).get("slides", []) if r["slide_id"] == canonical["slide_id"]), None)
        objects = slide["objects"]
        counts = Counter(o["kind"] for o in objects)
        family = exp["archetype_rendered"] if exp else canonical["visual_type"]
        families.append(family)
        body = [o for o in objects if not any(k in o["name"] for k in (":title", ":source", ":page-number", ":footer")) or o["kind"] == "image"]
        ratio = coverage([o["bounds"] for o in body]) / (13.333 * 7.5)
        # Reproducible anchor proxy: substantial non-text shapes/images/charts.
        anchors = sum(o["kind"] not in {"text", "line"} and o["bounds"]["w"] * o["bounds"]["h"] >= 0.5 for o in body)
        rows.append({"slide_index": index + 1, "slide_id": canonical["slide_id"], "composition_family": family,
                     "object_count": len(objects), "primitive_distribution": dict(counts), "visual_anchor_count": anchors,
                     "text_count": counts["text"], "coverage": ratio, "whitespace": 1 - ratio,
                     "text_density_characters_per_square_inch": sum(len(t) for t in slide["texts"]) / (13.333 * 7.5)})
    counts = Counter(o["kind"] for s in snapshot["slides"] for o in s["objects"])
    return {"metric_definition": "Bounding-box union coverage; native primitive counts; anchors are non-text body objects >=0.5 square inches. Descriptive only.",
            "layout_family_distribution": dict(Counter(s["layout_family"] for s in bundle["ir"]["slides"])),
            "rendered_composition_family_distribution": dict(Counter(families)), "primitive_distribution": dict(counts),
            "total_shape_count": sum(counts.values()), "roundRect_count": counts["roundRect"], "line_connector_count": counts["line"],
            "ellipse_count": counts["ellipse"], "text_object_count": counts["text"], "image_count": counts["image"], "chart_count": counts["chart"],
            "non_consecutive_visual_repetition": sum(families[i] == families[j] for i in range(len(families)) for j in range(i + 2, len(families))),
            "visual_anchor_count": sum(r["visual_anchor_count"] for r in rows),
            "mean_coverage": sum(r["coverage"] for r in rows) / len(rows),
            "mean_whitespace": sum(r["whitespace"] for r in rows) / len(rows),
            "mean_text_density": sum(r["text_density_characters_per_square_inch"] for r in rows) / len(rows), "slides": rows}


def run_arm(bundle: dict, arm: Path, enabled: bool) -> dict:
    start = time.perf_counter()
    pptx = generate(arm_spec(bundle, enabled), arm)
    generated = time.perf_counter()
    shutil.copyfile(bundle["staging"] / "deck_ir.json", arm / "deck_ir.json")
    geometry, log = inspect_powerpoint_text_layout(pptx_path=pptx, output_json=arm / "powerpoint_geometry.json",
        script_path=ROOT / "scripts/inspect_pptx_layout.ps1", timeout_seconds=180,
        preview_dir=arm / "preview/slides", output_pdf=arm / "deck.pdf")
    (arm / "powerpoint.log").write_text(log, encoding="utf-8")
    previews = sorted((arm / "preview/slides").glob("slide_*.png"))
    create_contact_sheet(previews, arm / "contact_sheet.png")
    ooxml = inspect_package(pptx, known_source_ids={r["source_id"] for r in bundle["manifest"]}, deck_ir=bundle["ir"])
    write_json(arm / "ooxml_qa.json", ooxml)
    write_json(arm / "object_manifest.json", ooxml["object_manifest"])
    rows = to_storyboard_rows(bundle["ir"])
    write_traceability_maps(arm, rows, bundle["claims"], [])
    bindings = [{"slide_id": s["slide_id"], "claim_ids": s["claim_ids"], "source_ids": s["source_ids"], "figure_ids": s["figure_ids"]} for s in bundle["ir"]["slides"]]
    known_claims = {c["claim_id"]: c for c in bundle["claims"]}
    known_sources = {s["source_id"] for s in bundle["manifest"]}
    binding_valid = all(set(s["source_ids"]) <= known_sources and s["claim_ids"] and all(
        cid in known_claims and known_claims[cid]["source_id"] in s["source_ids"]
        and known_claims[cid]["claim_text"] == s["key_message"] for cid in s["claim_ids"])
        for s in bundle["ir"]["slides"])
    write_json(arm / "source_claim_binding_qa.json", {"status": "PASS" if binding_valid else "FAIL", "slide_count": len(bindings)})
    write_json(arm / "source_bindings.json", bindings)
    with (arm / "source_binding_map.csv").open("x", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(bindings[0]))
        writer.writeheader()
        writer.writerows({k: ";".join(v) if isinstance(v, list) else v for k, v in row.items()} for row in bindings)
    planned_issues = [i for s in bundle["canonical"]["slides"] for i in validate_planned_geometry(s["planned_geometry"], bundle["contract"])]
    privacy_issues = []
    scanned, external = _scan_office_package(pptx, "deck.pptx", privacy_issues)
    privacy = {"status": "PASS" if not privacy_issues else "FAIL", "issues": privacy_issues, "office_xml_parts_scanned": scanned, "external_relationship_count": external}
    write_json(arm / "privacy_qa.json", privacy)
    qa = run_qa(pptx, arm / "deck.pdf", previews, arm / "preview/layouts", rows, bundle["claims"], bundle["manifest"],
                bundle["input_root"], arm, "PowerPoint COM", ooxml_report=ooxml, powerpoint_layout_report=geometry,
                density_issues=write_density_reports(arm, bundle["ir"]))
    snapshot = package_snapshot(pptx)
    execution_path = arm / "preview/experimental-execution.json"
    execution = json.loads(execution_path.read_text()) if execution_path.exists() else None
    metric = metrics(snapshot, bundle, execution)
    write_json(arm / "metrics.json", metric)
    runtime = {"total_seconds": time.perf_counter() - start, "generation_seconds": generated - start,
               "scope": "generation, PowerPoint export/actual geometry, contact sheet, existing QA, metrics; shared preparation/tests excluded"}
    write_json(arm / "runtime.json", runtime)
    qa_result = {"status": qa[0], "scientific_issues": qa[1], "visual_issues": qa[2], "file_issues": qa[3],
                 "planned_geometry_issues": planned_issues, "actual_geometry_issues": inspect_text_geometry(geometry), "privacy": privacy}
    write_json(arm / "qa.json", qa_result)
    gate(qa[0] == "READY_FOR_ASSISTED_USE" and not planned_issues and not privacy_issues and geometry["status"] == "PASS" and binding_valid,
         f"{arm.name} existing QA failed: {qa_result}")
    return dict(snapshot=snapshot, metrics=metric, execution=execution, runtime=runtime, qa=qa_result)


def invariance(bundle: dict, a: dict, b: dict, root: Path) -> dict:
    sa = json.loads((root / "arm_a/build_spec.json").read_text(encoding="utf-8"))
    sb = json.loads((root / "arm_b/build_spec.json").read_text(encoding="utf-8"))
    fields = {label: [s.get(key) for s in sa["slides"]] == [s.get(key) for s in sb["slides"]]
              for label, key in [("sources", "source_ids"), ("claims", "claim_ids"), ("citations", "short_source_label"),
                                ("content_hashes", "content_hash"), ("chart_contracts", "chart_data"), ("wording_boundary", "prohibited_overstatement")]}
    fields["source_bindings"] = sha(root / "arm_a/source_bindings.json") == sha(root / "arm_b/source_bindings.json")
    fields["canonical_deck_ir"] = sha(root / "arm_a/deck_ir.json") == sha(root / "arm_b/deck_ir.json") == sha(bundle["staging"] / "deck_ir.json")
    fields["canonical_hash"] = sa["deck_ir"]["canonical_hash"] == sb["deck_ir"]["canonical_hash"] == bundle["ir"]["canonical_hash"]
    for s in (sa, sb):
        s.pop("experiment", None)
        s.pop("visual_execution_plan", None)
    fields["canonical_build_spec"] = sa == sb == bundle["canonical"]
    fields["input_hashes"] = not verify_input_hashes(bundle["input_root"], bundle["manifest"]) and sha(FIXTURE) == sha(bundle["input_root"] / "fixture.json")
    fields["speaker_notes"] = a["snapshot"]["notes"] == b["snapshot"]["notes"]
    fields["visible_text_multiset"] = [s["tokens"] for s in a["snapshot"]["slides"]] == [s["tokens"] for s in b["snapshot"]["slides"]]
    # Exact paragraph multiset preserves complete statements/word order in this
    # fixture, supplementing the more permissive textbox-splitting token metric.
    fields["visible_scientific_strings"] = all(
        Counter(normalized(p) for t in x["texts"] for p in t.splitlines() if p.strip()) ==
        Counter(normalized(p) for t in y["texts"] for p in t.splitlines() if p.strip())
        for x, y in zip(a["snapshot"]["slides"], b["snapshot"]["slides"], strict=True))
    expected = [t for s in bundle["fixture"]["slides"][4]["takeaways"] for t in s.splitlines()]
    fields["all_six_editorial_statements"] = all(
        Counter(normalized(p) for t in arm["snapshot"]["slides"][4]["texts"] for p in t.splitlines())[normalized(unit)] == 1
        for arm in (a, b) for unit in expected) and len(expected) == 6
    for name, prefix in [("chart_values_and_workbooks", ("ppt/charts/", "ppt/embeddings/")), ("source_figure_bytes", ("ppt/media/",))]:
        pa = {n: h for n, h in a["snapshot"]["parts"].items() if n.startswith(prefix)}
        pb = {n: h for n, h in b["snapshot"]["parts"].items() if n.startswith(prefix)}
        fields[name] = bool(pa) and pa == pb
    fields["source_figure_original_embedded"] = all(sha(bundle["input_root"] / "figure.svg") in
        [h for n, h in arm["snapshot"]["parts"].items() if n.startswith("ppt/media/")] for arm in (a, b))
    fields["non_target_slide_xml"] = all(a["snapshot"]["parts"][f"ppt/slides/slide{i}.xml"] == b["snapshot"]["parts"][f"ppt/slides/slide{i}.xml"] for i in (3, 6))
    return {"status": "PASS" if all(fields.values()) else "FAIL", "passed": sum(fields.values()), "total": len(fields),
            "percent": 100 * sum(fields.values()) / len(fields), "fields": fields, "canonical_hash": bundle["ir"]["canonical_hash"]}


def run_tests(modules: list[str], output: Path) -> dict:
    stream = io.StringIO()
    suite = unittest.defaultTestLoader.loadTestsFromNames(modules)
    count = suite.countTestCases()
    result = unittest.TextTestRunner(stream=stream, verbosity=2).run(suite)
    output.write_text(stream.getvalue(), encoding="utf-8")
    return {"command": f"{sys.executable} -B -m unittest " + " ".join(modules) + " -v", "discovered": count,
            "passed": result.testsRun - len(result.failures) - len(result.errors) - len(result.skipped),
            "failed": len(result.failures) + len(result.errors), "skipped": len(result.skipped), "status": "PASS" if result.wasSuccessful() else "FAIL"}


def execute(root: Path) -> None:
    root.mkdir(parents=True, exist_ok=False)
    audit = root / "audit/visual_fidelity_001"
    audit.mkdir(parents=True)
    test_results = {}
    for label, modules in [("existing_baseline", BASELINE_MODULES), ("preservation", PRESERVATION_MODULES),
                           ("new", ["tests.test_visual_fidelity_experiment_001"])]:
        print(f"Testing {label}", flush=True)
        test_results[label] = run_tests(modules, audit / f"{label}_tests.log")
        write_json(audit / f"{label}_tests.json", test_results[label])
        gate(test_results[label]["status"] == "PASS", f"{label} regression failed")
    bundle = prepare(root)
    print("Rendering arm_a", flush=True)
    a = run_arm(bundle, root / "arm_a", False)
    print("Rendering arm_b", flush=True)
    b = run_arm(bundle, root / "arm_b", True)
    inv = invariance(bundle, a, b, root)
    write_json(audit / "scientific_invariance.json", inv)
    gate(inv["status"] == "PASS", "EXPERIMENT_FAILED_SCIENTIFIC_INVARIANCE")
    fidelity = b["execution"]
    gate(fidelity and fidelity["targeted_slide_count"] == fidelity["rendered_slide_count"] == 4
         and fidelity["generic_fallback_count"] == 0
         and all(r["archetype_requested"] == r["archetype_rendered"] for r in fidelity["slides"]), "Execution fidelity failed")
    write_json(audit / "execution_fidelity.json", fidelity)
    write_json(audit / "geometry_comparison.json", {"arm_a": a["qa"], "arm_b": b["qa"]})
    runtime = {"arm_a": a["runtime"], "arm_b": b["runtime"], "ratio_b_over_a": b["runtime"]["total_seconds"] / a["runtime"]["total_seconds"], "blocking_threshold": None}
    write_json(audit / "runtime_comparison.json", runtime)
    for label, arm in [("arm_a", a), ("arm_b", b)]:
        shutil.copyfile(root / label / "contact_sheet.png", audit / f"{label}_contact_sheet.png")
        write_json(audit / f"{label}_metrics.json", arm["metrics"])
    target_indexes = (0, 1, 3, 4)
    editability = {"native_text_and_shapes": all(s["objects"] and all(o["kind"] not in {"image", "chart", "other"} for o in s["objects"]) for i, s in enumerate(b["snapshot"]["slides"]) if i in target_indexes),
                   "editable_chart_with_workbook": any(n.startswith("ppt/embeddings/") for n in b["snapshot"]["parts"]) and b["metrics"]["chart_count"] == 1,
                   "replaceable_original_image": inv["fields"]["source_figure_original_embedded"] and b["metrics"]["image_count"] == 1}
    write_json(audit / "editability.json", editability)
    gate(all(editability.values()), "Editability structural check failed")
    # Reviewer receives only this neutral folder + blank form. The technical
    # arm mapping is intentionally separate; this is not a claimed blinded trial.
    review = audit / "human_review"
    review.mkdir()
    variants = ["Variant 1", "Variant 2"]
    arms = ["arm_a", "arm_b"]
    secrets.SystemRandom().shuffle(arms)
    for variant, arm in zip(variants, arms, strict=True):
        shutil.copyfile(root / arm / "contact_sheet.png", review / f"{variant}.png")
        shutil.copyfile(root / arm / "deck.pdf", review / f"{variant}.pdf")
    fields = ["variant", "semantic_fit", "information_hierarchy", "visual_coherence", "visual_richness", "readability", "manual_fix_burden", "comments"]
    with (audit / "human_blind_review.csv").open("x", newline="", encoding="utf-8-sig") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields)
        writer.writeheader()
        writer.writerows({"variant": v} for v in variants)
    shutil.copyfile(audit / "human_blind_review.csv", review / "human_blind_review.csv")
    write_json(audit / "review_mapping_for_auditor.json", dict(zip(variants, arms, strict=True)))
    manifest = ["# Experimental implementation manifest", "", f"Baseline: {BASELINE}",
                "Only the exact ID plus boolean enabled=true loads the isolated renderer.",
                "Disabled/absent/ID-only flags retain current production dispatch. Enabled with wrong ID fails closed (negative test requirement).",
                "No canonical schema, core, config, package metadata, source extraction, scientific QA or public runner changed.",
                "Baseline renderer snapshot was compared by unit tests at the OOXML part level; it is a temporary test oracle, not another production renderer.",
                "Geometry-only ellipse endpoint clipping references v2.5; no legacy runtime/derivation reused.",
                "All texts, relationships and assignments are explicit fixture values; no semantic inference.", "", "Changed files:", *["- " + p for p in ALLOWED]]
    (audit / "implementation_manifest.md").write_text("\n".join(manifest) + "\n", encoding="utf-8")
    summary = ["# Synthetic experiment summary", "", "VISUAL_FIDELITY_EXPERIMENT_001_SYNTHETIC_COMPLETE",
               "Recommendation: PROCEED_TO_HUMAN_AB_REVIEW", "", "Automated gates passed. This is not aesthetic approval or production adoption.",
               f"Scientific invariance: {inv['passed']}/{inv['total']} (100%). Four intended compositions; zero fallback.",
               f"Runtime A={runtime['arm_a']['total_seconds']:.3f}s, B={runtime['arm_b']['total_seconds']:.3f}s, B/A={runtime['ratio_b_over_a']:.3f}; descriptive only.",
               "Control stores six editorial statements as three pairs of complete paragraphs, avoiding the frozen renderer's three-container limit.",
               "Scope is limited to short English synthetic text, four-node graphs, unlabeled edges, and existing takeaway frames. Other targets fail closed.",
               "Source-figure control is a fixture-authored SVG, embedded as a replaceable image; no evidence visual was generated by AI.",
               "Visible-text equality includes token multisets AND exact paragraph multisets, plus explicit six-statement checks.",
               "Embedded workbook comparison preserves every member except creation/modified timestamp values and ZIP metadata; chart XML and source-figure bytes are exact.",
               "PowerPoint COM was used for both arms; no claim of LibreOffice/cross-host visual equivalence is made.",
               "Only the human_review folder should be supplied for neutral scoring. The technical package reveals arm assignment.",
               "Human scores remain empty. Real-project validation, long text/CJK, labelled-edge layouts, approved assets, promotion and release are outside scope."]
    (audit / "experiment_summary.md").write_text("\n".join(summary) + "\n", encoding="utf-8")
    print(json.dumps({"status": "VISUAL_FIDELITY_EXPERIMENT_001_SYNTHETIC_COMPLETE", "root": str(root), "tests": test_results, "invariance": inv, "runtime": runtime}, indent=2), flush=True)


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output-root", type=Path, required=True, help="New isolated synthetic experiment directory")
    args = parser.parse_args()
    gate(sys.dont_write_bytecode, "Run with Python -B; no bytecode may be written")
    try:
        execute(args.output_root.resolve())
    except RuntimeError as error:
        audit = args.output_root.resolve() / "audit/visual_fidelity_001"
        audit.mkdir(parents=True, exist_ok=True)
        write_json(audit / "failure.json", {"status": "VISUAL_FIDELITY_EXPERIMENT_001_BLOCKED", "error": str(error)})
        raise
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
