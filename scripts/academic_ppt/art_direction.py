"""Apply explicitly approved native style operations while preserving scientific content."""
from pathlib import Path
import hashlib
import copy
import zipfile
from pptx import Presentation
from .deck_ir import canonical_json_hash

def sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()

def gate(condition, message):
    if not condition:
        raise RuntimeError(message)

def apply_native_art_direction(source: Path, output: Path, plan: dict) -> dict:
    """Style a hash-bound native deck copy with explicit operations.

    Serialize only targeted slide XML (and explicitly styled synthetic charts).
    The existing PptxGenJS deck, notes, masters, media, bindings and untouched
    slides remain the authority. This function never re-derives scientific semantics.
    """
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
    from pptx.enum.text import PP_ALIGN
    from pptx.util import Inches, Pt
    from lxml import etree

    gate(plan.get("composition") == "NATIVE_ART_DIRECTION" and plan.get("enabled") is True,
         "ART_DIRECTION_COMPOSITION_NOT_ENABLED")
    gate(source.is_file() and sha(source) == plan.get("source_pptx_sha256"), "Frozen style source mismatch")
    gate(not output.exists() and source.resolve() != output.resolve(), "Refusing to overwrite a source or prior artifact")
    gate(canonical_json_hash(plan.get("semantic_plan")) == plan.get("semantic_plan_hash"), "Semantic plan hash mismatch")
    gate(plan.get("font_family") in {"Microsoft YaHei", "Arial"}, "Use a validated installed font")
    palette = set(plan["palette"].values())
    gate(palette and all(isinstance(c,str) and len(c)==6 and all(v in "0123456789ABCDEFabcdef" for v in c) for c in palette), "Invalid approved palette")
    allowed_style = {"bounds","font_size_pt","bold","color","align","line_spacing","fill","stroke","line_width_pt","geometry","runs"}
    deck = Presentation(source)
    changed_parts, records = {}, []
    seen_pages = set()

    def color(value):
        gate(value in palette, "Unapproved color")
        return RGBColor.from_string(value)

    def native_text(shape):
        # python-pptx's text accessor creates txBody on an empty autoshape.
        # Inspect XML presence first, so even a read cannot mutate footer lines.
        body=shape._element.find("{http://schemas.openxmlformats.org/presentationml/2006/main}txBody")
        return shape.text if body is not None else ""

    def bounds(shape, values):
        gate(len(values)==4 and all(isinstance(v,(int,float)) for v in values), "Invalid native bounds")
        x,y,w,h=values
        gate(x>=0 and y>=0 and w>=0 and h>=0 and x+w<=13.3331 and y+h<=6.6375, "Style geometry exceeds safe frame/footer")
        shape.left,shape.top,shape.width,shape.height = (Inches(v) for v in values)

    def style_font(font, values):
        font.name=plan["font_family"]
        if "font_size_pt" in values:
            gate(values["font_size_pt"]>=18, "Style would shrink core text below minimum")
            font.size=Pt(values["font_size_pt"])
        if "bold" in values:font.bold=bool(values["bold"])
        if "color" in values:font.color.rgb=color(values["color"])

    def chart_semantics(xml):
        tree=copy.deepcopy(xml)
        # Only these style subtrees are written by chart styling below. All
        # caches, axes/scales, category order and relationship IDs remain exact.
        for tag in ("spPr","txPr"):
            for node in list(tree.iter("{http://schemas.openxmlformats.org/drawingml/2006/chart}"+tag)):
                node.getparent().remove(node)
        return etree.tostring(tree,method="c14n")

    for target in plan["targets"]:
        page=target["page"]
        gate(page not in seen_pages and 1<=page<=len(deck.slides), "Duplicate or invalid style target")
        seen_pages.add(page)
        slide=deck.slides[page-1];sid=target["slide_id"]
        gate(sid in slide.notes_slide.notes_text_frame.text, "Style slide identity mismatch")
        original_text=[native_text(s) for s in slide.shapes if native_text(s)]
        shape_map={s.name:s for s in slide.shapes}
        gate(len(shape_map)==len(slide.shapes), "Ambiguous native shape names")
        slide.background.fill.solid();slide.background.fill.fore_color.rgb=color(target["background"])
        operations=[]
        for op in target["operations"]:
            name=op["shape_name"]
            gate(name in shape_map and not any(p in name for p in (":sources",":page-number",":footer")), "Unknown or protected style target")
            shape=shape_map[name];values=op["style"]
            gate(set(values)<=allowed_style, "Unsupported style operation")
            if ":edge-" in name:
                gate(set(values)<={"stroke","line_width_pt"} and values.get("stroke") in palette
                     and values.get("line_width_pt",1)>=0.5, "Semantic edge may only receive visible line styling")
            gate(not shape.has_chart and shape.shape_type!=13, "Scientific figure/chart requires explicit protected control")
            if native_text(shape):
                gate(op.get("expected_text")==shape.text, "Frozen visible text mismatch")
            if "bounds" in values:bounds(shape,values["bounds"])
            if "geometry" in values:
                gate(values["geometry"] in {"rect","roundRect","ellipse"} and not native_text(shape), "Unsupported shape grammar change")
                prst=shape._element.find(".//{http://schemas.openxmlformats.org/drawingml/2006/main}prstGeom")
                gate(prst is not None, "No existing native preset geometry")
                prst.set("prst",values["geometry"])
            if "fill" in values:
                if values["fill"] is None:shape.fill.background()
                else:shape.fill.solid();shape.fill.fore_color.rgb=color(values["fill"])
            if "stroke" in values:
                if values["stroke"] is None:shape.line.fill.background()
                else:shape.line.color.rgb=color(values["stroke"])
            if "line_width_pt" in values:shape.line.width=Pt(values["line_width_pt"])
            if native_text(shape):
                if "runs" in values:
                    gate(len(shape.text_frame.paragraphs)==1 and "".join(r["text"] for r in values["runs"])==shape.text,
                         "Styled runs must preserve exact text and order")
                    p=shape.text_frame.paragraphs[0];p.clear()
                    for token in values["runs"]:
                        run=p.add_run();run.text=token["text"];style_font(run.font,{**values,**token})
                else:
                    for p in shape.text_frame.paragraphs:
                        for run in p.runs:style_font(run.font,values)
                for p in shape.text_frame.paragraphs:
                    if "align" in values:p.alignment={"left":PP_ALIGN.LEFT,"center":PP_ALIGN.CENTER,"right":PP_ALIGN.RIGHT}[values["align"]]
                    if "line_spacing" in values:
                        gate(values["line_spacing"]>=1.05,"Line spacing below approved minimum")
                        p.line_spacing=values["line_spacing"]
            operations.append({"shape_name":name,"style":values})

        # Reclip only explicitly named existing edges to their declared nodes.
        # Roles and arrowheads are not inferred or changed by this style step.
        for edge in target.get("reclip_edges",[]):
            frozen=plan["semantic_plan"][sid]["edges"][edge["edge_index"]]
            gate(edge["source_shape"]==f"composition:{sid}:node-{frozen['source']}"
                 and edge["target_shape"]==f"composition:{sid}:node-{frozen['target']}"
                 and edge["shape_name"]==f"composition:{sid}:edge-{edge['edge_index']+1}", "Style cannot change semantic endpoints")
            shape=shape_map[edge["shape_name"]]
            a=shape_map[edge["source_shape"]];b=shape_map[edge["target_shape"]]
            ax,ay=(a.left+a.width/2)/914400,(a.top+a.height/2)/914400
            bx,by=(b.left+b.width/2)/914400,(b.top+b.height/2)/914400
            dx,dy=bx-ax,by-ay;length=(dx*dx+dy*dy)**0.5
            gate(length>0,"Coincident style nodes")
            ta=min(a.width/914400/2/abs(dx) if dx else float("inf"),a.height/914400/2/abs(dy) if dy else float("inf"))
            tb=min(b.width/914400/2/abs(dx) if dx else float("inf"),b.height/914400/2/abs(dy) if dy else float("inf"))
            gate(ta+tb+0.09/length<1,"Style nodes overlap")
            x1,y1=ax+dx*(ta+0.045/length),ay+dy*(ta+0.045/length)
            x2,y2=bx-dx*(tb+0.045/length),by-dy*(tb+0.045/length)
            bounds(shape,[min(x1,x2),min(y1,y2),abs(x2-x1),abs(y2-y1)])
            xfrm=shape._element.find(".//{http://schemas.openxmlformats.org/drawingml/2006/main}xfrm")
            for key,value in [("flipH",x2<x1),("flipV",y2<y1)]:
                if value:xfrm.set(key,"1")
                elif key in xfrm.attrib:del xfrm.attrib[key]

        for i,decor in enumerate(target.get("decorations",[])):
            gate(decor.get("role")=="presentation_only" and decor.get("kind") in {"line","rect"}, "Only non-evidence native ornament is allowed")
            if decor["kind"]=="line":
                x,y,w,h=decor["bounds"]
                shape=slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x),Inches(y),Inches(x+w),Inches(y+h))
                bounds(shape,decor["bounds"]);shape.line.color.rgb=color(decor["color"]);shape.line.width=Pt(decor.get("width_pt",1))
            else:
                shape=slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,0,0,1,1);bounds(shape,decor["bounds"])
                shape.fill.solid();shape.fill.fore_color.rgb=color(decor["color"]);shape.line.fill.background()
            shape.name=f"art:{sid}:presentation-only:{i+1}"
            if decor.get("behind"):
                slide.shapes._spTree.remove(shape._element);slide.shapes._spTree.insert(2,shape._element)

        for control in target.get("chart_controls",[]):
            gate(plan.get("synthetic_only") is True, "Chart style control must be synthetic")
            gate(set(control)=={"shape_name","series_color"}, "Unsupported chart treatment; axis/value changes forbidden")
            chart=shape_map[control["shape_name"]].chart
            before=chart_semantics(chart._chartSpace)
            for series in chart.series:series.format.fill.solid();series.format.fill.fore_color.rgb=color(control["series_color"])
            chart.font.name=plan["font_family"];chart.font.size=Pt(18)
            chart.category_axis.tick_labels.font.size=Pt(16);chart.value_axis.tick_labels.font.size=Pt(16)
            chart.value_axis.major_gridlines.format.line.color.rgb=color("E8E1D7")
            gate(before==chart_semantics(chart._chartSpace), "Chart values, scale, categories or other semantics changed")
            changed_parts[str(chart.part.partname).lstrip("/")]=etree.tostring(chart._chartSpace,xml_declaration=True,encoding="UTF-8",standalone=True)
        gate(original_text==[native_text(s) for s in slide.shapes if native_text(s)], "ART_DIRECTION_SCIENTIFIC_MUTATION")
        changed_parts[f"ppt/slides/slide{page}.xml"]=etree.tostring(slide._element,xml_declaration=True,encoding="UTF-8",standalone=True)
        records.append({"page":page,"slide_id":sid,"narrative_role":target["narrative_role"],"background":target["background"],"operations":operations,
                        "decorations":target.get("decorations",[]),"visible_text_exact":True})
    output.parent.mkdir(parents=True,exist_ok=True)
    with zipfile.ZipFile(source) as z,zipfile.ZipFile(output,"x",zipfile.ZIP_DEFLATED) as result:
        gate(set(changed_parts)<=set(z.namelist()), "Style pass introduced an unexpected package part")
        for info in z.infolist():result.writestr(info,changed_parts.get(info.filename,z.read(info.filename)))
    return {"status":"PASS","plan_hash":canonical_json_hash(plan),"source_sha256":sha(source),"output_sha256":sha(output),
            "changed_parts":sorted(changed_parts),"slides":records,"semantic_plan_hash":plan["semantic_plan_hash"],"production_code_changed":False}
