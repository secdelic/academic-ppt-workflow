from __future__ import annotations

import csv
import re
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET


def _xml_text(xml_bytes: bytes) -> str:
    root = ET.fromstring(xml_bytes)
    return "\n".join(
        text.strip() for text in root.itertext() if text and text.strip()
    )


def extract_source(path: Path) -> tuple[str, int | str, list[str]]:
    suffix = path.suffix.lower()
    warnings: list[str] = []
    if suffix in {".md", ".txt"}:
        return path.read_text(encoding="utf-8-sig", errors="replace"), 1, warnings
    if suffix in {".csv", ".tsv"}:
        delimiter = "\t" if suffix == ".tsv" else ","
        rows: list[str] = []
        with path.open("r", encoding="utf-8-sig", errors="replace", newline="") as handle:
            for row in csv.reader(handle, delimiter=delimiter):
                rows.append(" | ".join(row))
        return "\n".join(rows), 1, warnings
    if suffix == ".docx":
        from docx import Document

        document = Document(path)
        parts = [p.text for p in document.paragraphs if p.text.strip()]
        for table_index, table in enumerate(document.tables, start=1):
            parts.append(f"[TABLE {table_index}]")
            parts.extend(" | ".join(cell.text for cell in row.cells) for row in table.rows)
        return "\n".join(parts), len(document.sections), warnings
    if suffix == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(str(path))
        pages: list[str] = []
        for index, page in enumerate(reader.pages, start=1):
            text = page.extract_text() or ""
            pages.append(f"[PAGE {index}]\n{text}")
        joined = "\n".join(pages)
        if len(re.sub(r"\s+", "", joined)) < max(40, len(reader.pages) * 20):
            warnings.append(
                "LIKELY_SCANNED_OR_TEXT_EXTRACTION_FAILED; OCR not run; OCR source and uncertainty required"
            )
        return joined, len(reader.pages), warnings
    if suffix == ".pptx":
        parts: list[str] = []
        with zipfile.ZipFile(path) as archive:
            slides = sorted(
                name
                for name in archive.namelist()
                if re.fullmatch(r"ppt/slides/slide\d+\.xml", name)
            )
            for index, name in enumerate(slides, start=1):
                parts.append(f"[SLIDE {index}]\n{_xml_text(archive.read(name))}")
            notes = sorted(
                name
                for name in archive.namelist()
                if re.fullmatch(r"ppt/notesSlides/notesSlide\d+\.xml", name)
            )
            for index, name in enumerate(notes, start=1):
                parts.append(f"[NOTES {index}]\n{_xml_text(archive.read(name))}")
        return "\n".join(parts), len(slides), warnings
    if suffix == ".xlsx":
        from openpyxl import load_workbook

        workbook = load_workbook(path, read_only=True, data_only=False)
        parts: list[str] = []
        for sheet in workbook.worksheets:
            parts.append(f"[SHEET {sheet.title}]")
            for row in sheet.iter_rows(values_only=True):
                values = ["" if value is None else str(value) for value in row]
                if any(values):
                    parts.append(" | ".join(values))
        count = len(workbook.worksheets)
        workbook.close()
        return "\n".join(parts), count, warnings
    if suffix in {".png", ".jpg", ".jpeg"}:
        from PIL import Image

        with Image.open(path) as image:
            text = f"[IMAGE] format={image.format}; width={image.width}; height={image.height}; mode={image.mode}"
        warnings.append("Image content not OCR-extracted; manual figure review required")
        return text, 1, warnings
    if suffix == ".svg":
        text = _xml_text(path.read_bytes())
        return f"[SVG]\n{text}", 1, warnings
    raise ValueError(f"Unsupported source type: {suffix}")


def extract_all(input_root: Path, staging_root: Path, manifest: list[dict[str, str]]) -> dict[str, str]:
    extracted: dict[str, str] = {}
    extracted_dir = staging_root / "extracted"
    extracted_dir.mkdir(parents=True, exist_ok=True)
    for row in manifest:
        path = input_root / row["relative_path"]
        try:
            text, count, warnings = extract_source(path)
            row["parsed_successfully"] = "yes"
            row["page_or_sheet_count"] = str(count)
            row["parse_warning"] = "; ".join(warnings)
            extracted[row["source_id"]] = text
            (extracted_dir / f"{row['source_id']}.txt").write_text(text, encoding="utf-8")
        except Exception as exc:
            row["parsed_successfully"] = "no"
            row["parse_warning"] = f"{type(exc).__name__}: {exc}"
            extracted[row["source_id"]] = ""
    return extracted
