"""Explicit human decisions, immutable image hashes and slide-specific use authorization."""
from __future__ import annotations
import argparse,csv,json,re,zipfile,copy
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches
from .manual_image_handoff import check,read_json,sha_file,write_new,now,within,external_root,DECISIONS
from .deck_ir import canonical_json_hash


def authorization_digest(asset):
    return canonical_json_hash({k:v for k,v in asset.items() if k!='authorization_sha256'})


def apply_approvals(import_root,decision_csv,approved_root,*,reviewer=None,enabled=False):
    if not enabled:return None
    check(isinstance(reviewer,str) and bool(reviewer.strip()),'EXPLICIT_HUMAN_REVIEWER_REQUIRED')
    root=Path(import_root).resolve();binding=read_json(root/'intake_binding.json');review=read_json(root/'review_manifest.json')
    check(review['approval_inferred'] is False and review['registry_sha256']==binding['registry_sha256']==sha_file(root/'candidate_asset_registry.json'),'REVIEW_REGISTRY_CHANGED')
    check(review['baseline_sha256']==binding['baseline']['sha256']==sha_file(binding['baseline']['path']),'REVIEW_BASELINE_CHANGED')
    candidates={x['candidate_id']:x for x in read_json(root/'candidate_asset_registry.json')['candidates']}
    previews={x['candidate_id']:x for x in review['entries']}
    with Path(decision_csv).open(encoding='utf-8-sig',newline='') as stream:decisions=list(csv.DictReader(stream))
    seen=set();approved=[];slides=set()
    for row in decisions:
        cid=row.get('candidate_id');check(cid in candidates and cid not in seen,'UNKNOWN_OR_DUPLICATE_CANDIDATE');seen.add(cid)
        c=candidates[cid];check(row.get('asset_request_id')==c['asset_request_id'] and row.get('target_slide_id')==c['target_slide_id'],'DECISION_SCOPE_MISMATCH')
        decision=row.get('decision','').strip();approve=row.get('approve','').strip().lower()
        check(decision in ['',*DECISIONS] and approve in ('','false','no','0','true','yes','1'),'INVALID_HUMAN_DECISION')
        if approve not in ('true','yes','1') or decision!='APPROVE':continue
        check(c['target_slide_id'] not in slides,'MULTIPLE_APPROVED_CANDIDATES_FOR_SLIDE');slides.add(c['target_slide_id'])
        check(c['approval_status']=='CANDIDATE' and c['human_approved'] is False and c['scientific_evidence'] is False and c['presentation_only'] is True,'CANDIDATE_BOUNDARY_CHANGED')
        source=within(root/'assets'/c['filename'],root/'assets');check(sha_file(source)==c['sha256'],'CANDIDATE_HASH_MISMATCH')
        preview=previews[cid];preview_path=within(root/'review'/preview['preview_filename'],root/'review')
        check(sha_file(preview_path)==preview['preview_sha256'] and preview['asset_sha256']==c['sha256'] and preview['target_slide_id']==c['target_slide_id'],'REVIEW_PREVIEW_CHANGED')
        approved.append((c,source,preview))
    if not approved:return {'status':'NO_APPROVAL_APPLIED','approved_count':0}
    destination=external_root(approved_root);check(not destination.exists(),'APPROVED_OUTPUT_ALREADY_EXISTS')
    # A new immutable approval batch; never overwrites an existing project registry.
    destination.mkdir(parents=True);assets=[]
    for c,source,preview in approved:
        filename=c['candidate_id']+source.suffix.lower();(destination/filename).write_bytes(source.read_bytes())
        assets.append({'asset_id':c['candidate_id'],'file':filename,'role':'PRESENTATION_ILLUSTRATION','source':'MANUAL_HUMAN_TRIGGER',
          'scientific_evidence':False,'presentation_only':True,'human_approved':True,'status':'APPROVED','approved_by':reviewer.strip(),'approved_at':now(),
          'sha256':c['sha256'],'asset_request_id':c['asset_request_id'],'slide_ids':[c['target_slide_id']],'slide_index':preview['slide_index'],
          'baseline_sha256':review['baseline_sha256'],'focal_region':preview['focal_region'],'allowed_transforms':['contain'],
          'reviewed_preview_sha256':preview['preview_sha256'],'decision_sha256':sha_file(decision_csv),'review_manifest_sha256':sha_file(root/'review_manifest.json')})
    for asset in assets:asset['authorization_sha256']=authorization_digest(asset)
    write_new(destination/'asset_registry.yaml',{'schema_version':'approved-presentation-assets/1','assets':assets})
    return {'status':'EXPLICIT_HUMAN_APPROVAL_APPLIED','approved_count':len(assets),'registry':str(destination/'asset_registry.yaml')}


def validate_use(asset,asset_root,*,slide_id,baseline_sha256):
    check(asset.get('status')=='APPROVED' and asset.get('human_approved') is True and asset.get('scientific_evidence') is False and asset.get('presentation_only') is True,'APPROVED_PRESENTATION_ASSET_REQUIRED')
    check(slide_id in asset.get('slide_ids',[]) and asset.get('baseline_sha256')==baseline_sha256,'ASSET_SLIDE_OR_BASELINE_SCOPE_MISMATCH')
    check(asset.get('approved_by') and asset.get('approved_at') and asset.get('decision_sha256') and asset.get('reviewed_preview_sha256'),'HUMAN_APPROVAL_RECEIPT_REQUIRED')
    check(asset.get('authorization_sha256')==authorization_digest(asset),'ASSET_AUTHORIZATION_CHANGED')
    path=within(Path(asset_root)/asset['file'],asset_root);check(sha_file(path)==asset.get('sha256'),'APPROVED_ASSET_HASH_MISMATCH')
    return path


def main():
    parser=argparse.ArgumentParser(description=__doc__);parser.add_argument('--apply-approvals',action='store_true')
    parser.add_argument('--import-root',type=Path);parser.add_argument('--decisions',type=Path);parser.add_argument('--approved-root',type=Path);parser.add_argument('--reviewer')
    parser.add_argument('--assemble',action='store_true');parser.add_argument('--source',type=Path);parser.add_argument('--output',type=Path)
    parser.add_argument('--registry',type=Path);parser.add_argument('--slide-ids',type=Path,help='JSON array of source-bound slide IDs in deck order')
    args=parser.parse_args()
    if args.assemble:
        if args.apply_approvals:parser.error('approval and assembly are separate human steps')
        if not all((args.source,args.output,args.registry,args.slide_ids)):parser.error('source, output, registry and slide IDs are required')
        report=assemble_approved_assets(args.source,args.output,args.registry,read_json(args.slide_ids),enabled=True)
        write_new(args.output.with_suffix('.assembly.json'),report)
        print(json.dumps(report));return
    if not args.apply_approvals:return
    if not all((args.import_root,args.decisions,args.approved_root,args.reviewer)):parser.error('import root, decisions, approved root and reviewer are required')
    print(json.dumps(apply_approvals(args.import_root,args.decisions,args.approved_root,reviewer=args.reviewer,enabled=True)))


def assemble_approved_assets(source,output,registry,slide_ids,*,enabled=False):
    """Add authorized images into reviewed empty slots, keeping all prior scientific parts."""
    if not enabled:return None
    import io
    from xml.etree import ElementTree as E
    from PIL import Image
    source,output,registry=Path(source).resolve(),Path(output).resolve(),Path(registry).resolve()
    check(source.is_file() and not output.exists() and source!=output,'NEW_ASSEMBLY_OUTPUT_REQUIRED')
    document=read_json(registry);check(document.get('schema_version')=='approved-presentation-assets/1','APPROVED_REGISTRY_REQUIRED')
    baseline_sha=sha_file(source);deck=Presentation(source)
    check(isinstance(slide_ids,list) and len(slide_ids)==len(deck.slides) and len(set(slide_ids))==len(slide_ids),'SLIDE_ID_BINDING_REQUIRED')
    assets=document['assets'];check(bool(assets),'APPROVED_ASSETS_REQUIRED');seen=set();changed=set();receipts=[]
    for asset in assets:
        index=asset['slide_index'];check(type(index) is int and 1<=index<=len(slide_ids) and index not in seen,'UNIQUE_ASSET_SLIDE_REQUIRED');seen.add(index)
        path=validate_use(asset,registry.parent,slide_id=slide_ids[index-1],baseline_sha256=baseline_sha)
        from .illustration_requests import rectangle
        region=asset['focal_region'];rectangle(region)
        x,y,w,h=region['x']*deck.slide_width,region['y']*deck.slide_height,region['w']*deck.slide_width,region['h']*deck.slide_height
        with Image.open(path) as im:iw,ih=im.size
        scale=min(w/iw,h/ih);pw,ph=iw*scale,ih*scale;px,py=x+(w-pw)/2,y+(h-ph)/2
        slide=deck.slides[index-1]
        for shape in slide.shapes:
            intersects=min(px+pw,shape.left+shape.width)>max(px,shape.left) and min(py+ph,shape.top+shape.height)>max(py,shape.top)
            # shape.text can create txBody on an empty producer-native shape.
            # Read XML directly so the collision audit cannot mutate the source.
            has_text=any((node.text or '').strip() for node in shape.element.iter('{http://schemas.openxmlformats.org/drawingml/2006/main}t'))
            background=not has_text and any(k in shape.name.lower() for k in ('background','surface','decoration'))
            check(not intersects or background,'APPROVED_ASSET_WOULD_OCCLUDE_NATIVE_CONTENT')
        picture=slide.shapes.add_picture(str(path),round(px),round(py),round(pw),round(ph));picture.name='approved-asset:'+asset['asset_id']
        changed.add(str(slide.part.partname).lstrip('/'));receipts.append({'asset_id':asset['asset_id'],'slide_id':slide_ids[index-1],'source_sha256':asset['sha256'],'placement_inches':[px/914400,py/914400,pw/914400,ph/914400],'scientific_evidence':False,'presentation_only':True})
    memory=io.BytesIO();deck.save(memory)
    with zipfile.ZipFile(source) as z:original={n:z.read(n) for n in z.namelist()}
    with zipfile.ZipFile(memory) as z:modified={n:z.read(n) for n in z.namelist()}
    parts=dict(original)
    for part in changed:
        parts[part]=modified[part];folder,name=part.rsplit('/',1);rel=folder+'/_rels/'+name+'.rels';parts[rel]=modified[rel]
        before=E.fromstring(original[part]);after=E.fromstring(parts[part])
        from .master_layout import tree,xml,texts
        check(texts(before)==texts(after),'ASSET_INSERTION_SCIENTIFIC_TEXT_CHANGED')
        check([xml(s) for s in tree(before)]==[xml(s) for s in list(tree(after))[:-1]],'ASSET_INSERTION_NATIVE_OBJECT_CHANGED')
    for name,data in modified.items():
        if name not in original:
            check(name.startswith('ppt/media/'),'UNEXPECTED_ASSET_PACKAGE_PART');parts[name]=data
    # Merge only required media content types; retain all original declarations.
    ct='http://schemas.openxmlformats.org/package/2006/content-types';types=E.fromstring(parts['[Content_Types].xml']);known={(x.tag,x.get('Extension'),x.get('PartName')) for x in types}
    for item in E.fromstring(modified['[Content_Types].xml']):
        if (item.tag,item.get('Extension'),item.get('PartName')) not in known:types.append(item)
    parts['[Content_Types].xml']=E.tostring(types,encoding='utf-8',xml_declaration=True)
    output.parent.mkdir(parents=True,exist_ok=True)
    with zipfile.ZipFile(output,'x',zipfile.ZIP_DEFLATED) as z:
        for name,data in parts.items():z.writestr(name,data)
    from .ooxml_qa import inspect_pptx_ooxml
    check(inspect_pptx_ooxml(output)['valid'],'APPROVED_ASSET_OOXML_INVALID')
    check(sha_file(source)==baseline_sha,'SOURCE_CHANGED_DURING_ASSEMBLY')
    return {'status':'APPROVED_ASSETS_ASSEMBLED','source_sha256':baseline_sha,'output_sha256':sha_file(output),'scientific_invariance':True,
      'registry':{'path':str(registry),'sha256':sha_file(registry)},'assets':receipts}

if __name__=='__main__':main()
