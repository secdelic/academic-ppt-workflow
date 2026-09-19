"""Local comparison previews; candidate artwork never becomes scientific evidence."""
from __future__ import annotations
import io,json
from pathlib import Path
from PIL import Image,ImageDraw,ImageOps
from pptx import Presentation
from pptx.util import Inches,Pt
from .manual_image_handoff import check,read_json,sha_file,write_new,within
from .rendering import render_pptx,create_contact_sheet


def create_preview(import_root,*,enabled=False):
    if not enabled:return None
    root=Path(import_root).resolve();binding=read_json(root/'intake_binding.json')
    check(sha_file(root/'candidate_asset_registry.json')==binding['registry_sha256'],'REGISTRY_HASH_MISMATCH')
    candidates=read_json(root/'candidate_asset_registry.json')['candidates']
    if not candidates:return None
    baseline=binding['baseline'];check(sha_file(baseline['path'])==baseline['sha256'],'BASELINE_HASH_MISMATCH')
    preview=root/'review';check(not preview.exists(),'OUTPUT_ALREADY_EXISTS');preview.mkdir()
    method,slides,log=render_pptx(Path(baseline['path']),preview/'baseline.pdf',preview/'baseline',Path(__file__).resolve().parents[1],180)
    request_map={x['asset_request_id']:x for x in binding['packets']}
    deck=Presentation();deck.slide_width=Inches(16);deck.slide_height=Inches(9)
    comparison=[];receipts=[]
    for candidate in candidates:
        check(candidate['human_approved'] is False and candidate['approval_status']=='CANDIDATE' and candidate['scientific_evidence'] is False,'UNAPPROVED_CANDIDATE_REQUIRED')
        asset=within(root/'assets'/candidate['filename'],root/'assets');check(sha_file(asset)==candidate['sha256'],'CANDIDATE_HASH_MISMATCH')
        row=request_map[candidate['asset_request_id']];index=row['slide_index'];check(1<=index<=len(slides),'SLIDE_BINDING_INVALID')
        with Image.open(slides[index-1]) as image:base=image.convert('RGBA')
        with Image.open(asset) as image:art=image.convert('RGBA')
        # Show only the approved planning slot; do not invent full-background placement.
        region=row['request']['focal_region'];w,h=base.size
        box=(round(region['x']*w),round(region['y']*h),round(region['w']*w),round(region['h']*h))
        resized=ImageOps.contain(art,(box[2],box[3]));mockup=base.copy()
        mockup.alpha_composite(resized,(box[0]+(box[2]-resized.width)//2,box[1]+(box[3]-resized.height)//2))
        canvas=Image.new('RGB',(1600,1000),'#F6F2EA');draw=ImageDraw.Draw(canvas)
        draw.text((30,20),'CANDIDATE / NOT APPROVED / PRESENTATION ONLY',fill='#8E3B2F',font_size=28)
        draw.text((30,65),candidate['candidate_id']+' / '+candidate['target_slide_id'],fill='#233F39',font_size=20)
        left=ImageOps.contain(base.convert('RGB'),(740,740));right=ImageOps.contain(mockup.convert('RGB'),(740,740))
        canvas.paste(left,(30,160));canvas.paste(right,(830,160));draw.text((30,120),'BASELINE',fill='#233F39',font_size=22)
        draw.text((830,120),'LOCAL CANDIDATE MOCKUP',fill='#233F39',font_size=22)
        draw.text((30,925),'Local review only. Native scientific content remains authoritative. No approval inferred.',fill='#233F39',font_size=22)
        path=preview/(candidate['candidate_id']+'.png');canvas.save(path);comparison.append(path)
        slide=deck.slides.add_slide(deck.slide_layouts[6]);slide.shapes.add_picture(str(path),0,0,width=deck.slide_width,height=deck.slide_height)
        receipts.append({'candidate_id':candidate['candidate_id'],'asset_sha256':candidate['sha256'],'target_slide_id':candidate['target_slide_id'],
          'slide_index':index,'baseline_sha256':baseline['sha256'],'focal_region':region,'preview_filename':path.name,'preview_sha256':sha_file(path)})
    deck.save(preview/'candidate_comparison.pptx')
    images=[Image.open(p).convert('RGB') for p in comparison]
    try:images[0].save(preview/'candidate_comparison.pdf','PDF',save_all=True,append_images=images[1:],resolution=144)
    finally:
        for image in images:image.close()
    create_contact_sheet(comparison,preview/'candidate_contact_sheet.png')
    check(sha_file(baseline['path'])==baseline['sha256'],'BASELINE_CHANGED_DURING_PREVIEW')
    write_new(root/'review_manifest.json',{'status':'HUMAN_REVIEW_REQUIRED','registry_sha256':binding['registry_sha256'],
      'baseline_sha256':baseline['sha256'],'renderer':method,'entries':receipts,'approval_inferred':False,
      'preview_scientific_content_editability':'RASTER_LOCAL_COMPARISON_ONLY','source_pptx_modified':False})
    return preview
