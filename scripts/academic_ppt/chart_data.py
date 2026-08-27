from __future__ import annotations

import csv
import hashlib
import json
import re
from dataclasses import asdict, dataclass, field
from pathlib import Path
from typing import Any, Iterable, Mapping

from openpyxl import load_workbook


@dataclass(frozen=True)
class ChartDataContract:
    chart_id: str
    content_type: str
    source_id: str
    source_path: str
    categories: list[str]
    values: list[float]
    units: str
    numerator: list[float | None] = field(default_factory=list)
    denominator: list[float | None] = field(default_factory=list)
    sample_size_total: float | None = None
    sample_size_by_group: dict[str, float] = field(default_factory=dict)
    estimate_type: str = ""
    confidence_interval: list[dict[str, float]] = field(default_factory=list)
    reference_group: str = ""
    reference_value: float | None = None
    missingness: dict[str, float] = field(default_factory=dict)
    notes: list[str] = field(default_factory=list)
    canonical_status: str = "canonical"
    series: list[dict[str, Any]] = field(default_factory=list)

    def to_dict(self) -> dict[str, Any]:
        value = asdict(self)
        value["schema_version"] = "2.1"
        value["content_hash"] = hashlib.sha256(
            json.dumps(value, ensure_ascii=False, sort_keys=True).encode("utf-8")
        ).hexdigest()
        return value


def _float(value: Any) -> float | None:
    try:
        if value is None or str(value).strip() == "":
            return None
        return float(str(value).strip())
    except (TypeError, ValueError):
        return None


def _first(headers: Mapping[str, str], names: Iterable[str]) -> str | None:
    for name in names:
        if name in headers:
            return headers[name]
    return None


def _chart_id(source_id: str, content_type: str, discriminator: str = "") -> str:
    digest = hashlib.sha256(
        f"{source_id}:{content_type}:{discriminator}".encode("utf-8")
    ).hexdigest()
    return "CHT-" + digest[:16].upper()


def _read_table(path: Path) -> list[dict[str, str]]:
    delimiter = "\t" if path.suffix.lower() == ".tsv" else ","
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle, delimiter=delimiter))


def contract_from_records(
    *,
    source_id: str,
    source_path: str,
    records: list[dict[str, str]],
) -> ChartDataContract | None:
    """Infer a chart contract from an already-parsed aggregate table.

    The function is structural: it uses column contracts, never project names,
    page numbers, titles, or fixed category labels.
    """

    if not records:
        return None
    headers = {str(key).strip().lower(): str(key) for key in records[0]}
    category = _first(
        headers,
        ["phenotype", "category", "group", "contrast", "analysis", "subgroup", "variable", "time"],
    )
    if not category:
        return None

    estimate = _first(headers, ["estimate", "effect", "odds_ratio", "or", "hazard_ratio"])
    ci_low = _first(headers, ["ci_low", "lower_ci", "lower"])
    ci_high = _first(headers, ["ci_high", "upper_ci", "upper"])
    if estimate and ci_low and ci_high:
        rows = [
            row
            for row in records
            if _float(row.get(estimate)) is not None
            and _float(row.get(ci_low)) is not None
            and _float(row.get(ci_high)) is not None
        ]
        if rows:
            if "subgroup" in headers:
                content_type = "subgroup_results"
            elif "analysis" in headers:
                content_type = "sensitivity_analyses"
            else:
                content_type = "effect_estimates"
            categories = [str(row.get(category, "")).strip() for row in rows]
            values = [float(row[estimate]) for row in rows]
            intervals = [
                {"low": float(row[ci_low]), "high": float(row[ci_high])}
                for row in rows
            ]
            estimate_type_header = _first(headers, ["estimate_type", "measure"])
            estimate_type = (
                str(rows[0].get(estimate_type_header, "")).strip()
                if estimate_type_header
                else "effect estimate"
            )
            notes: list[str] = []
            interaction_header = _first(headers, ["interaction_p", "p_interaction"])
            if interaction_header:
                notes.append(
                    "interaction_p="
                    + ",".join(
                        str(row.get(interaction_header, "")).strip() for row in rows
                    )
                )
            return ChartDataContract(
                chart_id=_chart_id(source_id, content_type, source_path),
                content_type=content_type,
                source_id=source_id,
                source_path=source_path,
                categories=categories,
                values=values,
                units="ratio",
                estimate_type=estimate_type,
                confidence_interval=intervals,
                reference_value=1.0,
                notes=notes,
            )

    missing_n = _first(headers, ["missing_n", "missing_count"])
    missing_percent = _first(headers, ["missing_percent", "missing_rate_percent"])
    if missing_percent:
        rows = [
            row for row in records if _float(row.get(missing_percent)) is not None
        ]
        if rows:
            categories = [str(row.get(category, "")).strip() for row in rows]
            values = [float(row[missing_percent]) for row in rows]
            numerators = [
                _float(row.get(missing_n)) if missing_n else None for row in rows
            ]
            return ChartDataContract(
                chart_id=_chart_id(source_id, "missingness", source_path),
                content_type="missingness",
                source_id=source_id,
                source_path=source_path,
                categories=categories,
                values=values,
                units="percent",
                numerator=numerators,
                missingness=dict(zip(categories, values, strict=True)),
            )

    denominator = _first(
        headers,
        ["n", "denominator", "sample_size", "group_n", "total_n"],
    )
    rate = _first(
        headers,
        [
            "rate_percent",
            "event_rate_percent",
            "aki_rate_percent",
            "percentage",
            "percent",
        ],
    )
    if rate is None:
        rate = next(
            (original for lower, original in headers.items() if lower.endswith("_rate_percent")),
            None,
        )
    numerator = _first(headers, ["numerator", "event_n", "events"])
    if numerator is None:
        numerator = next(
            (
                original
                for lower, original in headers.items()
                if "event" in lower and not lower.endswith("rate_percent")
            ),
            None,
        )
    if denominator and rate:
        rows = [
            row
            for row in records
            if _float(row.get(denominator)) is not None
            and _float(row.get(rate)) is not None
        ]
        if rows:
            categories = [str(row.get(category, "")).strip() for row in rows]
            denominators = [float(row[denominator]) for row in rows]
            values = [float(row[rate]) for row in rows]
            numerators = [
                _float(row.get(numerator)) if numerator else None for row in rows
            ]
            return ChartDataContract(
                chart_id=_chart_id(
                    source_id, "categorical_event_rates", source_path
                ),
                content_type="categorical_event_rates",
                source_id=source_id,
                source_path=source_path,
                categories=categories,
                values=values,
                units="percent",
                numerator=numerators,
                denominator=denominators,
                sample_size_total=sum(denominators),
                sample_size_by_group=dict(
                    zip(categories, denominators, strict=True)
                ),
                notes=["Percentages are displayed from the canonical table; no recalculation is used."],
            )

    time_header = _first(headers, ["time", "timepoint", "visit"])
    value_header = _first(headers, ["value", "mean", "median"])
    series_header = _first(headers, ["series", "group", "phenotype"])
    if time_header and value_header and series_header:
        rows = [
            row for row in records if _float(row.get(value_header)) is not None
        ]
        if rows:
            series_names = list(
                dict.fromkeys(str(row[series_header]).strip() for row in rows)
            )
            time_labels = list(
                dict.fromkeys(str(row[time_header]).strip() for row in rows)
            )
            series = []
            for series_name in series_names:
                by_time = {
                    str(row[time_header]).strip(): float(row[value_header])
                    for row in rows
                    if str(row[series_header]).strip() == series_name
                }
                series.append(
                    {
                        "name": series_name,
                        "labels": time_labels,
                        "values": [by_time.get(label) for label in time_labels],
                    }
                )
            return ChartDataContract(
                chart_id=_chart_id(
                    source_id, "longitudinal_measurements", source_path
                ),
                content_type="longitudinal_measurements",
                source_id=source_id,
                source_path=source_path,
                categories=time_labels,
                values=series[0]["values"] if series else [],
                units="source unit",
                series=series,
                notes=[
                    "series="
                    + ",".join(series_names)
                ],
            )
    return None


def discover_chart_contracts(
    input_root: Path,
    manifest: list[dict[str, str]],
) -> list[dict[str, Any]]:
    contracts: list[ChartDataContract] = []
    for row in manifest:
        if row.get("source_role", "scientific_source") != "scientific_source":
            continue
        if row.get("file_type") not in {"csv", "tsv", "xlsx"}:
            continue
        path = input_root / row["relative_path"]
        if path.suffix.lower() == ".xlsx":
            workbook = load_workbook(path, read_only=True, data_only=True)
            try:
                for sheet in workbook.worksheets:
                    rows = list(sheet.iter_rows(values_only=True))
                    if not rows:
                        continue
                    headers = [
                        str(value).strip() if value is not None else ""
                        for value in rows[0]
                    ]
                    records = [
                        {
                            header: "" if value is None else str(value)
                            for header, value in zip(headers, values, strict=False)
                            if header
                        }
                        for values in rows[1:]
                        if any(value is not None for value in values)
                    ]
                    contract = contract_from_records(
                        source_id=row["source_id"],
                        source_path=f"{row['relative_path']}#{sheet.title}",
                        records=records,
                    )
                    if contract is not None:
                        contracts.append(contract)
            finally:
                workbook.close()
        else:
            records = _read_table(path)
            contract = contract_from_records(
                source_id=row["source_id"],
                source_path=row["relative_path"],
                records=records,
            )
            if contract is not None:
                contracts.append(contract)
    return [contract.to_dict() for contract in contracts]


def validate_chart_contract(contract: Mapping[str, Any]) -> list[str]:
    errors: list[str] = []
    required = [
        "chart_id", "source_id", "categories", "values", "units",
        "numerator", "denominator", "sample_size_total",
        "sample_size_by_group", "estimate_type", "confidence_interval",
        "reference_group", "missingness", "notes", "canonical_status",
    ]
    for field_name in required:
        if field_name not in contract:
            errors.append(f"missing field: {field_name}")
    categories = list(contract.get("categories", []))
    values = list(contract.get("values", []))
    if len(categories) != len(values):
        errors.append("categories and values have different lengths")
    denominators = list(contract.get("denominator", []))
    numerators = list(contract.get("numerator", []))
    if denominators and len(denominators) != len(categories):
        errors.append("denominator length does not match categories")
    if numerators and len(numerators) != len(categories):
        errors.append("numerator length does not match categories")
    intervals = list(contract.get("confidence_interval", []))
    if intervals and len(intervals) != len(categories):
        errors.append("confidence_interval length does not match categories")
    if denominators and any(value is None for value in denominators):
        errors.append("denominator contains missing values")
    return errors


__all__ = [
    "ChartDataContract",
    "contract_from_records",
    "discover_chart_contracts",
    "validate_chart_contract",
]
