"""Synthetic source-bound composition fixtures; independent of project state and Git history."""
from collections import Counter
import copy,csv,hashlib,io,json,os,re,shutil,subprocess,zipfile
from pathlib import Path
from xml.etree import ElementTree as ET
from pptx import Presentation
from scripts.academic_ppt.chart_data import ChartDataContract
from scripts.academic_ppt.deck_ir import build_deck_ir, canonical_json_hash, to_backend_spec
from scripts.academic_ppt.evidence import build_evidence
from scripts.academic_ppt.extractors import extract_all
from scripts.academic_ppt.inventory import inventory_sources, update_manifest
from scripts.academic_ppt.layout_contract import LayoutContract, plan_slide_geometry, validate_planned_geometry
from scripts.academic_ppt.art_direction import apply_native_art_direction
from scripts.academic_ppt.scientific_invariance import source_binding_issues as frozen_project_binding_issues
ROOT=Path(__file__).resolve().parents[2]
ID="SEMANTIC_COMPOSITION"
FIXTURE=ROOT/"tests/fixtures/synthetic/composition.json"
ZH_FIXTURE=ROOT/"tests/fixtures/synthetic/composition_zh.json"
D_FIXTURE=ROOT/"tests/fixtures/synthetic/semantic_grammar.json"
E_FIXTURE=ROOT/"tests/fixtures/synthetic/native_art_direction.json"
GENERATOR=ROOT/"scripts/generate_deck_pptxgen.mjs"
NS={"a":"http://schemas.openxmlformats.org/drawingml/2006/main"}

def write_json(path: Path, value) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("x", encoding="utf-8") as handle:
        json.dump(value, handle, ensure_ascii=False, indent=2)
        handle.write("\n")

def sha(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()

def gate(condition, message: str) -> None:
    if not condition:
        raise RuntimeError(message)

def prepare(root: Path, fixture_path: Path = FIXTURE) -> dict:
    """Register, extract, and freeze one synthetic source package for BOTH arms."""
    gate(fixture_path.resolve() in {FIXTURE.resolve(), ZH_FIXTURE.resolve(), D_FIXTURE.resolve(), E_FIXTURE.resolve()}, "Only registered synthetic fixtures are accepted")
    fixture = json.loads(fixture_path.read_text(encoding="utf-8"))
    gate(fixture["synthetic_only"] is True and fixture["fixture_id"] == ID, "Not the authorized synthetic fixture")
    source = root / "input"
    staging = root / "staging"
    source.mkdir(parents=True, exist_ok=False)
    staging.mkdir()
    shutil.copyfile(fixture_path, source / "fixture.json")
    (source / "figure.svg").write_text(fixture["source_figure_svg"], encoding="utf-8")
    with (source / "values.csv").open("x", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        writer.writerow(["category", "rate_percent"])
        writer.writerows(zip(fixture["chart"]["categories"], fixture["chart"]["values"], strict=True))
    lines = ["Fully synthetic software fixture. " + fixture["boundary"]]
    for slide in fixture["slides"]:
        lines.append(f"CLAIM|{slide['message']}|fixture {slide['key']}|synthetic software fixture|high|{slide['message']}|Do not infer clinical effects")
        lines.extend([slide["title"], *slide.get("takeaways", [])])
        if "diagram_spec" in slide:
            lines.append(json.dumps(slide["diagram_spec"], sort_keys=True))
    (source / "statements.md").write_text("\n".join(lines) + "\n", encoding="utf-8")
    manifest = inventory_sources(source, {".md", ".svg", ".csv"}, staging / "source_manifest.csv")
    extracted = extract_all(source, staging, manifest)
    update_manifest(staging / "source_manifest.csv", manifest)
    claims, unresolved = build_evidence(extracted, manifest, staging, fixture["brief"])
    gate(not unresolved, "Scientific source preparation failed: " + repr(unresolved))
    by_file = {r["file_name"]: r["source_id"] for r in manifest}
    contract = LayoutContract.from_files(ROOT / "config/layout_contract.yaml", ROOT / "config/safe_zones.yaml")
    chart = ChartDataContract(chart_id="CHT-SYNTHETIC-001", content_type="event_rates", source_id=by_file["values.csv"],
                              source_path="values.csv", **fixture["chart"]).to_dict()
    planned = []
    for index, f in enumerate(fixture["slides"]):
        s = {"logical_slide_key": f["key"], "slide_id": f["key"], "slide_role": f.get("slide_role","content"), "section_id": "synthetic",
             "slide_title": f["title"], "single_key_message": f["message"], "slide_purpose": "Synthetic composition validation",
             "visual_type": f["visual_type"], "layout_family": f["visual_type"], "proposed_layout": f["visual_type"],
             "source_ids": [by_file["statements.md"]], "claim_ids": [claims[index]["claim_id"]],
             "citation_requirement": "required", "short_source_label": "Source: synthetic fixture | software validation only",
             "speaker_note_summary": fixture["boundary"] + " " + f["message"], "prohibited_overstatement": "Do not infer clinical effects",
             "manual_review_required": "yes", "confidence": "high", "editable_object_requirements": ["native_text"]}
        for key in ("takeaways", "diagram_spec", "chart_caption"):
            if key in f:
                s[key] = copy.deepcopy(f[key])
        if f["visual_type"] == "editable_bar_chart":
            s.update(chart_data=chart, editable_object_requirements=["chart", "native_text"])
            s["source_ids"].append(by_file["values.csv"])
        if f["visual_type"] == "source_figure":
            s.update(visual_asset_path="input/figure.svg", figure_ids=["FIG-0001"], editable_object_requirements=["image", "native_text"])
            s["source_ids"].append(by_file["figure.svg"])
        s["planned_geometry"] = plan_slide_geometry(s, contract)
        gate(not validate_planned_geometry(s["planned_geometry"], contract), f"Planned geometry failed: {f['key']}")
        planned.append(s)
    ir = build_deck_ir(fixture["brief"], planned)
    canonical = to_backend_spec(ir, fixture["brief"], "restrained_biomedical")
    canonical["layout_contract"] = contract.serializable()
    if "style_profile" in fixture:
        canonical["style_profile"] = copy.deepcopy(fixture["style_profile"])
    assignments = {s["slide_id"]: copy.deepcopy(f["assignment"]) for s, f in zip(ir["slides"], fixture["slides"], strict=True) if "assignment" in f}
    for s in ir["slides"]:
        assignment = assignments.get(s["slide_id"], {})
        if assignment.get("revision") in {"COMPACT_CJK", "SEMANTIC"} and "diagram_spec" in s:
            assignment.update(copy.deepcopy(s["diagram_spec"]))
    write_json(staging / "deck_ir.json", ir)
    write_json(staging / "canonical_spec.json", canonical)
    write_json(staging / "freeze.json", {"fixture_sha256": sha(fixture_path), "canonical_hash": ir["canonical_hash"],
               "canonical_spec_sha256": sha(staging / "canonical_spec.json"), "source_hashes": {r["source_id"]: r["sha256"] for r in manifest}})
    return dict(fixture=fixture, manifest=manifest, claims=claims, ir=ir, canonical=canonical, assignments=assignments,
                input_root=source, staging=staging, contract=contract, fixture_path=fixture_path)

def arm_spec(bundle: dict, enabled: bool) -> dict:
    spec = copy.deepcopy(bundle["canonical"])
    spec["composition"] = {"id": ID, "enabled": enabled}
    if enabled:
        spec["visual_execution_plan"] = copy.deepcopy(bundle["assignments"])
    return spec

def generate(spec: dict, arm: Path, script: Path = GENERATOR) -> Path:
    arm.mkdir(parents=True, exist_ok=False)
    write_json(arm / "build_spec.json", spec)
    pptx = arm / "deck.pptx"
    node = shutil.which("node")
    gate(node, "Node.js is required")
    env = dict(os.environ, NODE_PATH=str(ROOT / "node_modules"), PYTHONDONTWRITEBYTECODE="1")
    result = subprocess.run([node, str(script), str(arm / "build_spec.json"), str(pptx), str(arm / "preview")],
                            cwd=arm.parent, env=env, capture_output=True, text=True, encoding="utf-8", errors="replace", timeout=120)
    (arm / "generation.log").write_text(result.stdout + result.stderr, encoding="utf-8")
    gate(result.returncode == 0, "Native generation failed: " + result.stderr)
    return pptx

def normalized_tokens(texts: list[str]) -> Counter:
    """Whitespace/box-order independent token multiset; retain punctuation/numbers.

    Paired with exact frozen-paragraph coverage below; token equality alone cannot
    prove that recommendations retain their word order or meaning.
    """
    return Counter(token for text in texts for token in re.findall(r"\w+|[^\w\s]", text, re.UNICODE))

def normalized(text: str) -> str:
    return " ".join(text.split())

def workbook_content_hash(data: bytes) -> str:
    # Office creation/modified timestamps and ZIP timestamps are nondeterministic
    # packaging metadata. Compare every decompressed workbook member, removing
    # ONLY the two core timestamp values; cell values/formulas/styles remain exact.
    with zipfile.ZipFile(io.BytesIO(data)) as workbook:
        members = {}
        for name in workbook.namelist():
            raw = workbook.read(name)
            if name == "docProps/core.xml":
                xml = ET.fromstring(raw)
                for tag in ("created", "modified"):
                    for element in xml.findall("{http://purl.org/dc/terms/}" + tag):
                        element.text = "TIMESTAMP"
                raw = ET.tostring(xml)
            members[name] = hashlib.sha256(raw).hexdigest()
    return canonical_json_hash(members)

def package_snapshot(path: Path) -> dict:
    deck = Presentation(path)
    slides = []
    for slide in deck.slides:
        texts, objects = [], []
        for shape in slide.shapes:
            text = shape.text if shape.has_text_frame else ""
            if text:
                texts.append(text)
            preset = shape._element.find(".//a:prstGeom", NS)
            kind = ("chart" if shape.has_chart else "image" if shape.shape_type == 13 else
                    "text" if text else preset.get("prst") if preset is not None else "other")
            if shape._element.tag.endswith("}cxnSp"):
                kind = "line"
            objects.append({"name": shape.name, "kind": kind, "text": text,
                            "bounds": {"x": shape.left / 914400, "y": shape.top / 914400,
                                       "w": shape.width / 914400, "h": shape.height / 914400}})
        slides.append({"texts": texts, "tokens": dict(normalized_tokens(texts)), "objects": objects})
    with zipfile.ZipFile(path) as z:
        names = z.namelist()
        parts = {n: hashlib.sha256(z.read(n)).hexdigest() for n in names if
                 re.fullmatch(r"ppt/(slides/slide\d+\.xml|notesSlides/notesSlide\d+\.xml|charts/.*|embeddings/.*|media/.*)", n)}
        for n in parts:
            if n.startswith("ppt/embeddings/") and n.endswith(".xlsx"):
                parts[n] = workbook_content_hash(z.read(n))
        notes = {n: [e.text or "" for e in ET.fromstring(z.read(n)).findall(".//a:t", NS)] for n in names if re.fullmatch(r"ppt/notesSlides/notesSlide\d+\.xml", n)}
    return {"slides": slides, "parts": parts, "notes": notes}
