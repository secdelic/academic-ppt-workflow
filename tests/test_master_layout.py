"""Real Master/Layout regressions on native synthetic decks."""
import copy,json,tempfile,unittest,zipfile
from pathlib import Path
from xml.etree import ElementTree as E
from pptx import Presentation
from pptx.util import Inches,Pt
from scripts.academic_ppt import master_layout as m


class MasterLayoutTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory(prefix='synthetic-master-');self.addCleanup(self.temp.cleanup);self.root=Path(self.temp.name)
        self.source=self.root/'source.pptx';self.output=self.root/'bound.pptx';deck=Presentation();deck.slide_width=Inches(13.333);deck.slide_height=Inches(7.5)
        self.bindings=[]
        for i,(role,layout) in enumerate(zip(('HERO','CONTEXT','EVIDENCE','FRAMEWORK','IMPLEMENTATION','TIMELINE','CLOSING'),m.LAYOUTS),1):
            slide=deck.slides.add_slide(deck.slide_layouts[6]);sid=f'S{i:03d}'
            for name,text,box in [('title',f'Synthetic role {i}',[.6,.4,11,.7]),('scientific',f'Value {i} unit; no causal claim',[.6,2,8,1]),('sources','Source: synthetic fixture',[.6,6.9,10,.25]),('page-number',str(i),[12,6.9,.6,.25])]:
                shape=slide.shapes.add_textbox(*(Inches(v) for v in box));shape.name=sid+':'+name;shape.text=text
                for run in shape.text_frame.paragraphs[0].runs:run.font.size=Pt(18)
            slide.notes_slide.notes_text_frame.text='[Sources]\nSynthetic fixture, immutable notes'
            self.bindings.append({'slide_id':sid,'slide_index':i,'page_role':role,'layout_id':layout,'binding_reason':'Synthetic explicit role','override_status':'NONE'})
        deck.save(self.source)

    def run_binding(self,**kwargs):return m.bind_master(self.source,self.output,self.bindings,enabled=True,**kwargs)

    def test_disabled_no_io(self):self.assertIsNone(m.bind_master(None,None,None))
    def test_real_master_layout_relationships(self):
        report=self.run_binding();self.assertEqual(len(report['layout_parts']),7)
        qa=m.inspect_pptx_ooxml(self.output);self.assertTrue(qa['valid'])
        for b,chain in zip(self.bindings,qa['slide_relationship_chains']):
            self.assertEqual(chain['layout_part'],report['layout_parts'][b['layout_id']]);self.assertEqual(chain['master_part'],report['master_part'])
    def test_no_scientific_text_in_shared_parts(self):
        report=self.run_binding()
        with zipfile.ZipFile(self.output) as z:
            for n in [report['master_part'],*report['layout_parts'].values()]:self.assertEqual(m.texts(E.fromstring(z.read(n))),[])
    def test_notes_and_scientific_bodies_unchanged(self):
        self.run_binding()
        with zipfile.ZipFile(self.source) as a,zipfile.ZipFile(self.output) as b:
            for name in a.namelist():
                if name.startswith('ppt/notesSlides/'):self.assertEqual(a.read(name),b.read(name))
            for i in range(1,8):
                n=f'ppt/slides/slide{i}.xml';old=E.fromstring(a.read(n));new=E.fromstring(b.read(n));self.assertEqual(m.texts(old),m.texts(new))
                scientific=lambda t:next(s for s in m.tree(t) if m.shape_name(s).endswith(':scientific'))
                self.assertEqual(m.xml(scientific(old)),m.xml(scientific(new)))
    def test_real_native_placeholders_inherit_font(self):
        report=self.run_binding();self.assertEqual(len(report['placeholder_bindings']),21)
        deck=Presentation(self.output)
        for slide in deck.slides:
            title=next(s for s in slide.shapes if s.name.endswith(':title'));self.assertTrue(title.is_placeholder)
            self.assertTrue(all(r.font.size is None for p in title.text_frame.paragraphs for r in p.runs))
    def test_footer_and_number_inheritance(self):
        report=self.run_binding()
        with zipfile.ZipFile(self.output) as z:
            self.assertIn('Master:footer-rule',z.read(report['master_part']).decode())
            for n in report['layout_parts'].values():self.assertIn('sldNum',z.read(n).decode())
    def test_layout_safe_zones_explicit(self):
        report=self.run_binding()
        for record in report['layouts'].values():
            for key in ('title_safe_zone','content_safe_zone','figure_safe_zone','footer_exclusion','negative_space_zone'):self.assertEqual(len(record[key]),4)
    def test_wrong_role_binding_fails(self):
        self.bindings[0]['layout_id']='EVIDENCE_CLEAN'
        with self.assertRaisesRegex(ValueError,'ROLE_MISMATCH'):self.run_binding()
        self.assertFalse(self.output.exists())
    def test_missing_binding_fails(self):
        with self.assertRaisesRegex(ValueError,'UNIQUE_MASTER_BINDINGS'):m.bind_master(self.source,self.output,[],enabled=True)
    def test_override_needs_receipt(self):
        b=self.bindings[0];b['override_status']='CUSTOM_SCENE_EXCEPTION';b['override_reason']='Synthetic explicit variation'
        with self.assertRaisesRegex(ValueError,'OVERRIDE_RECEIPT'):self.run_binding()
        b['override_receipt']={'slide_id':b['slide_id'],'approved_by':'SYNTHETIC_REVIEWER','reason':'Synthetic receipt'}
        self.assertEqual(self.run_binding()['bindings'][0]['master_binding_status'],'OVERRIDE_WITH_RECEIPT')
    def test_official_template_precedence_byte_identical(self):
        report=self.run_binding(official_template=True)
        self.assertEqual(self.source.read_bytes(),self.output.read_bytes());self.assertEqual(report['status'],'OFFICIAL_TEMPLATE_PRESERVED')
    def test_official_template_scene_insertion_refused(self):
        with self.assertRaisesRegex(ValueError,'OFFICIAL_TEMPLATE_ASSET'):self.run_binding(official_template=True,scene_by_layout={'HERO_SCENE':{}})
    def test_unapproved_master_art_rejected(self):
        with self.assertRaisesRegex(ValueError,'APPROVED_PRESENTATION_ASSET_REQUIRED'):self.run_binding(scene_by_layout={'HERO_SCENE':{'asset':{}}},asset_root=self.root)
        self.assertFalse(self.output.exists())
    def test_unknown_layout_rejected(self):
        self.bindings[0]['layout_id']='UNKNOWN'
        with self.assertRaisesRegex(ValueError,'UNKNOWN_MASTER_LAYOUT'):self.run_binding()
    def test_source_not_overwritten(self):
        with self.assertRaisesRegex(ValueError,'NEW_MASTER_OUTPUT_REQUIRED'):m.bind_master(self.source,self.source,self.bindings,enabled=True)
    def test_existing_output_not_overwritten(self):
        self.output.write_bytes(b'SYNTHETIC_SENTINEL')
        with self.assertRaises(ValueError):self.run_binding()
        self.assertEqual(self.output.read_bytes(),b'SYNTHETIC_SENTINEL')
    def test_invalid_canvas_fails_closed(self):
        deck=Presentation(self.source);deck.slide_width=Inches(10);deck.save(self.source)
        with self.assertRaisesRegex(ValueError,'WIDE_CANVAS_REQUIRED'):self.run_binding()
    def test_font_not_installed_not_silently_used(self):
        with self.assertRaisesRegex(ValueError,'VALIDATED_FONT_REQUIRED'):self.run_binding(theme={'head_font':'UNKNOWN_FONT'})
    def test_master_cannot_carry_page_specific_text(self):
        with self.assertRaisesRegex(ValueError,'MASTER_SCIENTIFIC_TEXT_FORBIDDEN'):self.run_binding(theme={'footer':'Synthetic claim not allowed'})
    def test_evidence_has_no_scene(self):
        report=self.run_binding()
        with zipfile.ZipFile(self.output) as z:self.assertNotIn('role-accent',z.read(report['layout_parts']['EVIDENCE_CLEAN']).decode())

    def test_master_and_layout_ids_globally_unique_for_powerpoint(self):
        self.run_binding();ids=[]
        with zipfile.ZipFile(self.output) as z:
            for n in z.namelist():
                if n=='ppt/presentation.xml' or (n.startswith('ppt/slideMasters/') and n.endswith('.xml')):
                    ids.extend(x.get('id') for x in E.fromstring(z.read(n)).iter() if x.tag in ('{'+m.P+'}sldMasterId','{'+m.P+'}sldLayoutId'))
        self.assertEqual(len(ids),len(set(ids)))
    def test_theme_roles_actually_materialized_without_changing_source_theme(self):
        report=self.run_binding();qa=m.inspect_pptx_ooxml(self.output);theme=qa['slide_relationship_chains'][0]['theme_part']
        with zipfile.ZipFile(self.source) as before,zipfile.ZipFile(self.output) as after:
            self.assertEqual(before.read('ppt/theme/theme1.xml'),after.read('ppt/theme/theme1.xml'))
            node=E.fromstring(after.read(theme));self.assertEqual(node.get('name'),m.DEFAULT_THEME['theme_id'])
            self.assertEqual(node.find('a:themeElements/a:fontScheme/a:majorFont/a:latin',m.NS).get('typeface'),'Microsoft YaHei')
    def test_charts_and_embedded_data_preserved(self):
        from pptx.chart.data import CategoryChartData
        from pptx.enum.chart import XL_CHART_TYPE
        deck=Presentation(self.source);data=CategoryChartData();data.categories=['Synthetic A','Synthetic B'];data.add_series('Synthetic values',[1,2])
        deck.slides[2].shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED,Inches(2),Inches(3),Inches(4),Inches(2),data);deck.save(self.source)
        self.run_binding()
        with zipfile.ZipFile(self.source) as before,zipfile.ZipFile(self.output) as after:
            parts=[n for n in before.namelist() if n.startswith(('ppt/charts/','ppt/embeddings/'))];self.assertTrue(parts)
            for name in parts:self.assertEqual(before.read(name),after.read(name))
        self.assertTrue(any(s.has_chart for s in Presentation(self.output).slides[2].shapes))

    def test_producer_reserved_content_type_names_do_not_collide(self):
        with zipfile.ZipFile(self.source) as z:parts={n:z.read(n) for n in z.namelist()}
        types=E.fromstring(parts['[Content_Types].xml'])
        for name,kind in [('slideMasters/slideMaster12.xml','slideMaster'),('slideLayouts/slideLayout30.xml','slideLayout')]:
            E.SubElement(types,'{'+m.CT+'}Override',{'PartName':'/ppt/'+name,'ContentType':f'application/vnd.openxmlformats-officedocument.presentationml.{kind}+xml'})
        parts['[Content_Types].xml']=m.xml(types)
        with zipfile.ZipFile(self.source,'w',zipfile.ZIP_DEFLATED) as z:
            for name,data in parts.items():z.writestr(name,data)
        report=self.run_binding()
        self.assertEqual(report['master_part'],'ppt/slideMasters/slideMaster13.xml')
        self.assertEqual(report['layout_parts']['HERO_SCENE'],'ppt/slideLayouts/slideLayout31.xml')
        self.assertTrue(m.inspect_pptx_ooxml(self.output)['valid'])

if __name__=='__main__':unittest.main()
