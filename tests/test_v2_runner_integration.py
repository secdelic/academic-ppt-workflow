from __future__ import annotations

import json
import shutil
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

import sys

from PIL import Image, ImageDraw
from pypdf import PdfWriter
from pptx import Presentation


ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))

from academic_ppt.qa import run_qa
from academic_ppt.runner import (
    STAGES,
    _build_object_manifest_document,
    execute,
)
from academic_ppt.routing import RouteValidationError
from academic_ppt.utils import sha256_file


def _write_repo_contract(root: Path) -> tuple[Path, Path]:
    (root / "config").mkdir(parents=True)
    (root / "brief").mkdir()
    (root / "input").mkdir()
    (root / "audit").mkdir()
    config = {
        "schema_version": "1.0",
        "paths": {
            "input_root": "input",
            "staging_root": "staging",  # self-containment: generated-temp
            "output_root": "output",
            "audit_root": "audit",
            "archive_root": "archive",
        },
        "environment": {
            "input_root": "TEST_UNUSED_INPUT",
            "staging_root": "TEST_UNUSED_STAGING",
            "output_root": "TEST_UNUSED_OUTPUT",
            "node_executable": "TEST_UNUSED_NODE",
            "node_modules": "TEST_UNUSED_NODE_MODULES",
            "soffice_executable": "TEST_UNUSED_SOFFICE",
            "pdftoppm_executable": "TEST_UNUSED_PDFTOPPM",
        },
        "supported_extensions": [".md", ".pptx", ".potx"],
        "default_backend": "pptxgenjs",
        "render_timeout_seconds": 30,
    }
    (root / "config" / "workflow.yaml").write_text(
        json.dumps(config), encoding="utf-8"
    )
    brief = {
        "project_name": "v2_runner_test",
        "presentation_type": "research_report",
        "presentation_objective": "Test the local route controller.",
        "target_slide_count": 6,
        "key_message": "All statements remain source-bound.",
        "style_reference_files": [],
    }
    brief_path = root / "brief" / "presentation_brief.yaml"
    brief_path.write_text(json.dumps(brief), encoding="utf-8")
    source = root / "input" / "study.md"
    source.write_text(
        "CLAIM|Synthetic exact claim.|section 1|reported fact|high|"
        "Synthetic exact claim.|Do not add causality\n",
        encoding="utf-8",
    )
    return brief_path, source


def _write_pptx(path: Path, slide_count: int = 1) -> None:
    presentation = Presentation()
    for index in range(slide_count):
        slide = presentation.slides.add_slide(presentation.slide_layouts[5])
        slide.shapes.title.text = f"Synthetic slide {index + 1}"
    path.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(path)


def _fake_generate(
    backend,
    node,
    node_modules,
    scripts_root,
    spec,
    pptx,
    artifact_preview,
    timeout,
):
    value = json.loads(Path(spec).read_text(encoding="utf-8"))
    _write_pptx(Path(pptx), len(value["slides"]))
    return f"FAKE_BACKEND={backend}"


def _fake_render(
    pptx_path,
    pdf_path,
    preview_dir,
    scripts_root,
    timeout,
    soffice,
    pdftoppm,
):
    slide_count = len(Presentation(pptx_path).slides)
    writer = PdfWriter()
    preview_dir.mkdir(parents=True, exist_ok=True)
    previews = []
    for index in range(slide_count):
        writer.add_blank_page(width=960, height=540)
        preview = preview_dir / f"slide_{index + 1:03d}.png"
        image = Image.new("RGB", (1600, 900), "white")
        draw = ImageDraw.Draw(image)
        draw.rectangle((0, 0, 799, 899), fill="#E7F0F7")
        draw.text((100, 100), f"Slide {index + 1}", fill="black")
        image.save(preview)
        previews.append(preview)
    with Path(pdf_path).open("wb") as handle:
        writer.write(handle)
    return "synthetic-test-renderer", previews, "synthetic render"


def _fake_ooxml(path, known_source_ids=None, deck_ir=None):
    identities = []
    objects = []
    for index, slide in enumerate(deck_ir["slides"], start=1):
        identities.append(
            {
                "slide_id": slide["slide_id"],
                "slide_index": index,
                "ooxml_canonical_sha256": f"{index:064x}",
            }
        )
        objects.append(
            {
                "slide_index": index,
                "slide_part": f"ppt/slides/slide{index}.xml",
                "manifest_sha256": f"{index + 100:064x}",
                "object_kind_counts": {"shape": 1},
                "object_count": 1,
                "chart_count": 0,
                "table_count": 0,
            }
        )
    return {
        "status": "PASS",
        "valid": True,
        "errors": [],
        "warnings": [],
        "source": {"sha256": sha256_file(Path(path))},
        "object_manifest": objects,
        "ir_identity_check": {
            "consistent": True,
            "identity_map": identities,
            "errors": [],
            "warnings": [],
        },
    }


def _args(root: Path, brief: Path, **overrides):
    values = {
        "brief": str(brief),
        "repo_root": str(root),
        "input_root": str(root / "input"),
        "staging_root": str(root / "staging"),
        "output_root": str(root / "output"),
        "run_id": "20260728_120000",
        "route": "generate",
        "backend": None,
        "node": None,
        "node_modules": None,
        "soffice": None,
        "pdftoppm": None,
        "from_stage": "preflight",
        "resume": False,
        "reference_pptx": None,
        "style_profile": None,
        "reference_mode": "style-only",
        "existing_pptx": None,
        "previous_deck_ir": None,
        "previous_object_manifest": None,
        "update_slide": None,
        "update_section": None,
        "update_figure": None,
        "allow_managed_overwrite": False,
        "template_layout_id": None,
        "narrative_mode": None,
    }
    values.update(overrides)
    return SimpleNamespace(**values)


class RunnerIntegrationTests(unittest.TestCase):
    def test_ooxml_qa_stage_precedes_visual_and_package(self):
        self.assertLess(STAGES.index("render"), STAGES.index("ooxml_qa"))
        self.assertLess(STAGES.index("ooxml_qa"), STAGES.index("visual_qa"))
        self.assertLess(STAGES.index("ooxml_qa"), STAGES.index("package"))

    def test_create_style_profile_route_is_read_only_and_atomic(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            brief, _ = _write_repo_contract(root)
            reference = root / "reference.pptx"
            _write_pptx(reference)
            before = sha256_file(reference)
            output = execute(
                _args(
                    root,
                    brief,
                    route="create-style-profile",
                    reference_pptx=str(reference),
                )
            )
            self.assertEqual(sha256_file(reference), before)
            profile = json.loads(
                (output / "reference_style_profile.json").read_text(
                    encoding="utf-8"
                )
            )
            runtime = json.loads(
                (output / "runtime_manifest.json").read_text(encoding="utf-8")
            )
            self.assertFalse(profile["privacy_contract"]["slide_text_returned"])
            self.assertEqual(runtime["route"], "create-style-profile")
            self.assertFalse(runtime["network_access"])

    def test_generate_route_keeps_pptxgenjs_default_and_packages_v2_audit(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            brief, _ = _write_repo_contract(root)
            with (
                patch("academic_ppt.runner._detect_node", return_value="node"),
                patch("academic_ppt.runner._run_generator", side_effect=_fake_generate),
                patch("academic_ppt.runner.render_pptx", side_effect=_fake_render),
                patch("academic_ppt.runner.inspect_package", side_effect=_fake_ooxml),
            ):
                output = execute(_args(root, brief))
            runtime = json.loads(
                (output / "runtime_manifest.json").read_text(encoding="utf-8")
            )
            self.assertEqual(runtime["route"], "generate")
            self.assertEqual(runtime["backend"], "pptxgenjs")
            self.assertEqual(runtime["status"], "READY_FOR_ASSISTED_USE")
            for name in (
                "deck_ir.json",
                "ooxml_qa_report.json",
                "object_manifest.json",
                "slide_manifest.csv",
                "claim_source_map.csv",
            ):
                self.assertTrue((output / name).is_file(), name)

    def test_fill_template_dispatches_native_backend_and_packages_style_profile(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            brief, _ = _write_repo_contract(root)
            template = root / "template.pptx"
            _write_pptx(template)

            def fake_fill(**kwargs):
                deck = json.loads(
                    Path(kwargs["deck_ir_path"]).read_text(encoding="utf-8")
                )
                _write_pptx(Path(kwargs["output_pptx"]), len(deck["slides"]))
                return "synthetic native fill"

            with (
                patch(
                    "academic_ppt.runner.run_native_template_fill",
                    side_effect=fake_fill,
                ) as native_fill,
                patch("academic_ppt.runner._run_generator") as stable_generator,
                patch("academic_ppt.runner.render_pptx", side_effect=_fake_render),
                patch("academic_ppt.runner.inspect_package", side_effect=_fake_ooxml),
            ):
                output = execute(
                    _args(
                        root,
                        brief,
                        route="fill-template",
                        reference_pptx=str(template),
                    )
                )
            self.assertTrue(native_fill.called)
            stable_generator.assert_not_called()
            runtime = json.loads(
                (output / "runtime_manifest.json").read_text(encoding="utf-8")
            )
            self.assertEqual(runtime["backend"], "native-template-fill")
            self.assertTrue((output / "reference_style_profile.json").is_file())

    def test_enhance_existing_keeps_source_and_packages_slide_diff(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            brief, _ = _write_repo_contract(root)
            common_patches = (
                patch("academic_ppt.runner._detect_node", return_value="node"),
                patch("academic_ppt.runner._run_generator", side_effect=_fake_generate),
                patch("academic_ppt.runner.render_pptx", side_effect=_fake_render),
                patch("academic_ppt.runner.inspect_package", side_effect=_fake_ooxml),
            )
            with common_patches[0], common_patches[1], common_patches[2], common_patches[3]:
                original_output = execute(_args(root, brief))

            existing = original_output / "v2_runner_test.pptx"
            existing_hash = sha256_file(existing)
            previous_ir_path = original_output / "deck_ir.json"
            previous_ir = json.loads(previous_ir_path.read_text(encoding="utf-8"))
            target_id = next(
                slide["slide_id"]
                for slide in previous_ir["slides"]
                if slide["logical_slide_key"] == "conclusion-author-gate"
            )
            brief_value = json.loads(brief.read_text(encoding="utf-8"))
            brief_value["key_message"] = "Updated source-bound conclusion."
            brief.write_text(json.dumps(brief_value), encoding="utf-8")

            def fake_replace(**kwargs):
                shutil.copy2(kwargs["existing_pptx"], kwargs["output_pptx"])
                return "synthetic incremental replacement"

            with (
                patch("academic_ppt.runner._detect_node", return_value="node"),
                patch("academic_ppt.runner._run_generator", side_effect=_fake_generate),
                patch(
                    "academic_ppt.runner._pptx_slide_id_order",
                    return_value=[
                        slide["slide_id"] for slide in previous_ir["slides"]
                    ],
                ),
                patch(
                    "academic_ppt.runner.run_powerpoint_replacement",
                    side_effect=fake_replace,
                ) as replacement,
                patch("academic_ppt.runner.render_pptx", side_effect=_fake_render),
                patch("academic_ppt.runner.inspect_package", side_effect=_fake_ooxml),
            ):
                updated_output = execute(
                    _args(
                        root,
                        brief,
                        run_id="20260728_120001",
                        route="enhance-existing",
                        workflow_mode="full_validation",
                        existing_pptx=str(existing),
                        previous_deck_ir=str(previous_ir_path),
                        previous_object_manifest=str(
                            original_output / "object_manifest.json"
                        ),
                        update_slide=target_id,
                    )
                )
            self.assertTrue(replacement.called)
            self.assertEqual(sha256_file(existing), existing_hash)
            diff = json.loads(
                (updated_output / "slide_diff.json").read_text(encoding="utf-8")
            )
            self.assertEqual(diff["targets"][0]["slide_id"], target_id)
            self.assertTrue(diff["targets"][0]["content_changed"])
            runtime = json.loads(
                (updated_output / "runtime_manifest.json").read_text(
                    encoding="utf-8"
                )
            )
            self.assertTrue(runtime["incremental_update"]["source_unchanged"])

    def test_generate_route_rejects_reference_as_a_second_top_level_route(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            brief, _ = _write_repo_contract(root)
            reference = root / "reference.pptx"
            _write_pptx(reference)
            with self.assertRaises(RouteValidationError):
                execute(
                    _args(
                        root,
                        brief,
                        route="generate",
                        reference_pptx=str(reference),
                    )
                )
            self.assertFalse((root / "output").exists())

    def test_ooxml_errors_gate_file_qa_and_are_reported(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            pptx_path = root / "deck.pptx"
            pdf_path = root / "deck.pdf"
            preview = root / "slide_001.png"
            output = root / "out"
            input_root = root / "input"
            output.mkdir()
            input_root.mkdir()
            _write_pptx(pptx_path)
            writer = PdfWriter()
            writer.add_blank_page(width=960, height=540)
            with pdf_path.open("wb") as handle:
                writer.write(handle)
            image = Image.new("RGB", (1600, 900), "white")
            ImageDraw.Draw(image).rectangle((0, 0, 799, 899), fill="blue")
            image.save(preview)
            status, _, _, file_issues = run_qa(
                pptx_path,
                pdf_path,
                [preview],
                root / "layouts",
                [
                    {
                        "slide_id": "SLD-X",
                        "citation_requirement": "none",
                        "source_ids": "",
                        "single_key_message": "Synthetic",
                    }
                ],
                [],
                [],
                input_root,
                output,
                "synthetic",
                ooxml_report={
                    "errors": [
                        {
                            "code": "DANGLING_RELATIONSHIP",
                            "message": "Synthetic relationship failure.",
                        }
                    ],
                    "warnings": [],
                },
            )
            self.assertEqual(status, "PARTIALLY_READY")
            self.assertTrue(
                any("DANGLING_RELATIONSHIP" in item for item in file_issues)
            )
            self.assertIn(
                "## OOXML QA",
                (output / "qa_report.md").read_text(encoding="utf-8"),
            )

    def test_object_manifest_uses_stable_slide_identity(self):
        report = {
            "source": {"sha256": "a" * 64},
            "object_manifest": [
                {
                    "slide_index": 1,
                    "slide_part": "ppt/slides/slide1.xml",
                    "manifest_sha256": "b" * 64,
                    "object_kind_counts": {"shape": 1},
                    "object_count": 1,
                }
            ],
            "ir_identity_check": {
                "identity_map": [{"slide_index": 1, "slide_id": "SLD-STABLE"}]
            },
        }
        document = _build_object_manifest_document(
            report,
            {
                "deck_id": "DECK-X",
                "deck_version": 2,
                "canonical_hash": "c" * 64,
            },
        )
        self.assertEqual(document["slides"][0]["slide_id"], "SLD-STABLE")
        self.assertEqual(document["slides"][0]["managed_object_hash"], "b" * 64)


if __name__ == "__main__":
    unittest.main()
