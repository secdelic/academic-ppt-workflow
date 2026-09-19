"""Synthetic manual image governance; no real project files or external asset directories."""
import copy,csv,io,json,os,re,shutil,subprocess,sys,tempfile,unittest
from pathlib import Path
from unittest.mock import patch
from PIL import Image
from pptx import Presentation
from pptx.util import Inches
from pptx.enum.shapes import MSO_SHAPE
from scripts.academic_ppt import manual_image_handoff as m
from scripts.academic_ppt import illustration_requests as ir
from scripts.academic_ppt.candidate_preview import create_preview
from tests.support.illustration_fixture import context_for,SUBJECTS
INITIAL=(1,3,6,14)


def synthetic_render(pptx,pdf,preview,*args,**kwargs):
    preview.mkdir(parents=True,exist_ok=True);paths=[]
    for i in range(len(Presentation(pptx).slides)):
        p=preview/f"slide_{i+1:03d}.png";Image.new('RGB',(640,360),'white').save(p);paths.append(p)
    pdf.write_bytes(b'SYNTHETIC_RENDER_TEST_DOUBLE')
    return 'SYNTHETIC_TEST_DOUBLE',paths,''


class ManualHandoffTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory(prefix='synthetic-handoff-');self.addCleanup(self.temp.cleanup)
        self.root=Path(self.temp.name);self.base=self.root/'external';self.base.mkdir()
        self.run='SYNTHETIC_TEST';self.package=self.root/'package';self.requests={}
        for index in SUBJECTS:
            record={'slide_id':f'S{index:03d}','page_role':'HERO','visual_anchor_role':'ANCHOR','recommended_ambition_level':'STRETCH',
              'enrichment_opportunity':'HIGH','quality_floor_status':'HUMAN_REVIEW','deck_quality_floor_status':'VISUAL_QUALITY_BLOCKED',
              'source_priority':copy.deepcopy(ir.SOURCE_PRIORITY),'source_figure_required':False,
              'recommended_next_action':'HUMAN_SEMANTIC_DECISION' if index in (5,9) else 'REFERENCE_GRADE_SCENE_TRIAL',
              'illustration_decision':'NO_ILLUSTRATION_NEEDED'}
            self.requests[index]=ir.build_request(record,context_for(index),{'approved':True,'style_family':ir.STYLE,'palette':ir.PALETTE},
              {'scientific_source_figure':False,'approved_project_asset':False,'approved_deterministic_library':False},enabled=True)
        deck=Presentation();deck.slide_width=Inches(13.333);deck.slide_height=Inches(7.5)
        for i in range(15):
            slide=deck.slides.add_slide(deck.slide_layouts[6]);slide.shapes.add_textbox(Inches(.5),Inches(.5),Inches(4),Inches(1)).text=f'Synthetic source {i+1}'
            # Some native generators emit empty vector shapes without txBody.
            mark=slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(.1),Inches(.1),Inches(.1),Inches(.1))
            for node in list(mark.element):
                if node.tag.endswith('}txBody'):mark.element.remove(node)
            slide.notes_slide.notes_text_frame.text='[Sources] Synthetic software fixture only'
        self.baseline=self.root/'source.pptx';deck.save(self.baseline)
        m.prepare(self.package,self.run,records=[{'slide_index':i,'request':self.requests[i]} for i in INITIAL],asset_root=self.base,baseline_pptx=self.baseline,enabled=True)
        self.inbox=m.authorized_run(self.run,self.base);self.manifest=m.read_json(self.package/'manual_handoff_manifest.json')
        self.out=self.base/'candidate_runs'/self.run/'IMPORT_TEST'
        self.render_patch=patch('scripts.academic_ppt.candidate_preview.render_pptx',side_effect=synthetic_render)
        self.render_patch.start();self.addCleanup(self.render_patch.stop)

    def candidate(self,index=1,slot='A',suffix='.png',color=(40,90,70,255)):
            req=self.requests[index]['request_id'];p=self.inbox/req/(req+'__'+slot+suffix)
            image=Image.new('RGBA',(48,27),color)
            if suffix in ('.jpg','.jpeg'):image.convert('RGB').save(p,'JPEG')
            else:image.save(p,'WEBP' if suffix=='.webp' else 'PNG')
            return p

    def intake(self):
            return m.intake(self.package,self.out,asset_root=self.base,enabled=True,import_image_candidates=True)

    def packet(self,index=1):
            return m.make_packet(self.requests[index],enabled=True)

    def test_default_off_has_no_project_io(self):
            with patch.object(Path,'read_bytes',side_effect=AssertionError),patch.object(Path,'read_text',side_effect=AssertionError),patch.object(Path,'mkdir',side_effect=AssertionError),patch.object(Path,'iterdir',side_effect=AssertionError):
                self.assertIsNone(m.make_packet(None));self.assertIsNone(m.prepare(None,None));self.assertIsNone(m.intake(None,None));self.assertIsNone(create_preview(None))

    def test_no_import_flag_does_not_scan(self):
            with patch.object(Path,'iterdir',side_effect=AssertionError),patch.object(m,'load_package',side_effect=AssertionError):
                self.assertIsNone(m.intake(None,None,enabled=True))

    def test_import_flag_without_enabled_does_nothing(self):
            with patch.object(m,'load_package',side_effect=AssertionError):
                self.assertIsNone(m.intake(None,None,import_image_candidates=True))

    def test_python_cli_default_off(self):
            cmd=[sys.executable,'-B','-m','scripts.academic_ppt.manual_image_handoff','--import-image-candidates','--package','MISSING','--output',str(self.out)]
            r=subprocess.run(cmd,capture_output=True)
            self.assertEqual((r.returncode,r.stdout,r.stderr),(0,b'',b''));self.assertFalse(self.out.exists())

    def test_unsafe_free_text_input(self):
            r=copy.deepcopy(self.requests[1]);r['free_text']='arbitrary source'
            with self.assertRaisesRegex(m.BridgeError,'MANUAL_GENERATION_PACKET_UNSAFE'):m.make_packet(r,enabled=True)

    def test_patient_identifier_leakage(self):
            r=copy.deepcopy(self.requests[1]);r['subject_summary']='SYNTHETIC_PATIENT_ID_001'
            with self.assertRaises(m.BridgeError):m.make_packet(r,enabled=True)

    def test_local_path_leakage_in_packet(self):
            packet=self.packet();packet['prompt']+=' C:'+chr(92)+'synthetic'+chr(92)+'private.txt'
            with self.assertRaisesRegex(m.BridgeError,'MANUAL_GENERATION_PACKET_UNSAFE'):m.validate_packet(packet,self.requests[1])

    def test_raw_manuscript_leakage(self):
            r=copy.deepcopy(self.requests[1]);r['manuscript']='SYNTHETIC_PRIVATE_MANUSCRIPT_PARAGRAPH'
            with self.assertRaises(m.BridgeError):m.make_packet(r,enabled=True)

    def test_prohibited_medical_request(self):
            r=copy.deepcopy(self.requests[6]);r['allowed_visual_types']=['EEG','MRI']
            with self.assertRaises(m.BridgeError):m.make_packet(r,enabled=True)

    def test_scientific_evidence_request_rejected(self):
            r=copy.deepcopy(self.requests[1]);r['scientific_evidence']=True
            with self.assertRaises(m.BridgeError):m.make_packet(r,enabled=True)

    def test_numeric_no_illustration_control(self):
            self.assertIsNone(m.make_packet(self.requests[11],enabled=True))

    def test_unresolved_no_packet(self):
            for index in [5,9]:self.assertIsNone(m.make_packet(self.requests[index],enabled=True))

    def test_required_scientific_boundary(self):
            for index in INITIAL:self.assertIn(m.BOUNDARY,self.packet(index)['prompt'])

    def test_prompt_structure_no_json(self):
            prompt=self.packet()['prompt']
            titles=['PURPOSE','SUBJECT','STYLE','PALETTE','COMPOSITION','NEGATIVE SPACE','VISUAL DEPTH','ALLOWED ELEMENTS','FORBIDDEN ELEMENTS','SCIENTIFIC BOUNDARY','OUTPUT PREFERENCE']
            self.assertEqual([part.splitlines()[0] for part in prompt.split('\n\n')],titles)
            self.assertNotIn('{',prompt);self.assertNotIn('asset_request_id',prompt)

    def test_palette_and_subject_are_controlled_values(self):
            for i in INITIAL:
                p=self.packet(i);self.assertEqual(p['palette'],self.requests[i]['palette']);self.assertEqual(p['short_conceptual_subject'],self.requests[i]['subject_summary'])
                for key in ('style_family','allowed_visual_types','prohibited_visual_types'):
                    self.assertEqual(p[key],self.requests[i][key])

    def test_markdown_contains_same_prompt_and_metadata(self):
            for i in INITIAL:
                p=self.packet(i);md=m.packet_markdown(p,i)
                self.assertEqual(md.split('```text\n')[1].split('\n```')[0],p['prompt'])
                for key in p:
                    if key!='prompt':self.assertIn('**'+key+'**:',md)
                self.assertNotRegex(md,r'[A-Za-z]:[\\/]')

    def test_asset_types_not_global(self):
            self.assertEqual({self.packet(i)['preferred_asset_type'] for i in INITIAL},{'FULL_SCENE','ICON_SET','TRANSPARENT_FOREGROUND','BACKGROUND'})

    def test_packet_deterministic_and_read_only(self):
            original=copy.deepcopy(self.requests);self.assertEqual(self.packet(),self.packet());self.assertEqual(original,self.requests)

    def test_no_private_paths_in_packets(self):
            for i in INITIAL:
                payload=json.dumps(self.packet(i));self.assertNotRegex(payload,r'[A-Za-z]:[\\/]');self.assertNotIn('source_registry',payload);self.assertNotIn('claim_registry',payload)

    def test_unknown_asset_request_ID(self):
            (self.inbox/'IRQ-UNKNOWN').mkdir()
            with self.assertRaisesRegex(m.BridgeError,'UNKNOWN_ASSET_REQUEST_ID'):self.intake()
            self.assertFalse(self.out.exists())

    def test_unknown_request_in_filename(self):
            p=self.candidate();p.rename(p.parent/'UNKNOWN__A.png')
            with self.assertRaisesRegex(m.BridgeError,'UNKNOWN_ASSET_REQUEST_ID_OR_FILENAME'):self.intake()

    def test_candidate_path_escape(self):
            p=self.root/'outside.png';Image.new('RGB',(4,4)).save(p)
            with self.assertRaisesRegex(m.BridgeError,'CANDIDATE_PATH_ESCAPE'):m.inspect_image(p,self.inbox,self.requests[1]['request_id'])

    def test_reparse_point_rejected(self):
            p=self.candidate();original=Path.lstat
            def attrs(path):
                result=original(path)
                if path==p:
                    from types import SimpleNamespace
                    return SimpleNamespace(st_mode=result.st_mode,st_file_attributes=0x400)
                return result
            with patch.object(Path,'lstat',attrs):
                with self.assertRaisesRegex(m.BridgeError,'CANDIDATE_PATH_ESCAPE'):m.no_reparse(p)

    def test_hardlink_rejected(self):
            p=self.candidate();os.link(p,self.root/'hardlink.png')
            with self.assertRaisesRegex(m.BridgeError,'CANDIDATE_PATH_ESCAPE'):self.intake()

    def test_nested_directory_rejected(self):
            req=self.requests[1]['request_id'];(self.inbox/req/(req+'__A.png')).mkdir()
            with self.assertRaises(m.BridgeError):self.intake()

    def test_unsupported_file(self):
            req=self.requests[1]['request_id'];(self.inbox/req/(req+'__A.svg')).write_text('<svg/>')
            with self.assertRaisesRegex(m.BridgeError,'UNSUPPORTED_FILE'):self.intake()

    def test_corrupt_file(self):
            req=self.requests[1]['request_id'];(self.inbox/req/(req+'__A.png')).write_bytes(b'not an image')
            with self.assertRaisesRegex(m.BridgeError,'CORRUPT_OR_UNSAFE_IMAGE'):self.intake()

    def test_mime_mismatch(self):
            p=self.candidate(suffix='.jpg');p.rename(p.with_suffix('.png'))
            with self.assertRaisesRegex(m.BridgeError,'MIME_EXTENSION_MISMATCH'):self.intake()

    def test_animated_file_rejected(self):
            p=self.candidate();a=Image.new('RGB',(5,5),'red');b=Image.new('RGB',(5,5),'blue');a.save(p,save_all=True,append_images=[b],duration=100,loop=0)
            with self.assertRaisesRegex(m.BridgeError,'ANIMATED_FILE_UNSUPPORTED'):self.intake()

    def test_pixel_budget(self):
            self.candidate()
            with patch.object(m,'MAX_PIXELS',100):
                with self.assertRaisesRegex(m.BridgeError,'INVALID_IMAGE_DIMENSIONS'):self.intake()

    def test_accept_one_candidate(self):
            self.candidate();self.intake();self.assertEqual(len(m.read_json(self.out/'candidate_asset_registry.json')['candidates']),1)

    def test_accept_two_candidates(self):
            self.candidate();self.candidate(slot='B',color=(100,10,50,255));self.intake();self.assertEqual(len(m.read_json(self.out/'candidate_asset_registry.json')['candidates']),2)

    def test_accept_three_candidates(self):
            for i,slot in enumerate('ABC'):self.candidate(slot=slot,color=(i*70,50,80,255))
            self.intake();self.assertEqual(len(m.read_json(self.out/'candidate_asset_registry.json')['candidates']),3)

    def test_more_than_three_rejected(self):
            for i,slot in enumerate('ABCD'):self.candidate(slot=slot,color=(i*60,50,80,255))
            with self.assertRaisesRegex(m.BridgeError,'TOO_MANY_CANDIDATES'):self.intake()

    def test_duplicate_slot_rejected(self):
            self.candidate();self.candidate(suffix='.jpg')
            with self.assertRaisesRegex(m.BridgeError,'DUPLICATE_CANDIDATE_SLOT'):self.intake()

    def test_candidate_hash_mime_dimensions_recorded(self):
            p=self.candidate();h=m.sha_file(p);self.intake();c=m.read_json(self.out/'candidate_asset_registry.json')['candidates'][0]
            self.assertEqual((c['sha256'],c['mime_type'],c['width'],c['height']),(h,'image/png',48,27))
            self.assertEqual(m.sha_file(self.out/'assets'/p.name),h);self.assertEqual(m.sha_file(p),h)

    def test_unapproved_status_and_file_presence_not_approval(self):
            self.candidate();self.intake();c=m.read_json(self.out/'candidate_asset_registry.json')['candidates'][0]
            self.assertEqual(c['approval_status'],'CANDIDATE');self.assertIs(c['human_approved'],False);self.assertIs(c['scientific_evidence'],False)
            self.assertEqual(c['generation_method'],'MANUAL_HUMAN_TRIGGER');self.assertEqual(c['generator_source'],'USER_DECLARED');self.assertIsNone(c['generator_identity'])

    def test_approval_table_blank(self):
            self.candidate();self.intake()
            with (self.out/'candidate_review_template.csv').open(encoding='utf-8-sig',newline='') as f: rows=list(csv.DictReader(f))
            self.assertEqual(len(rows),1);self.assertTrue(all(rows[0][k]=='' for k in ['approve','decision','comments']))

    def test_creation_and_import_times_recorded(self):
            self.candidate();self.intake();c=m.read_json(self.out/'candidate_asset_registry.json')['candidates'][0]
            self.assertTrue(c['imported_at']);self.assertTrue(c['file_creation_time']);self.assertIn('NOT_GENERATION_TIME',c['creation_time_basis'])

    def test_empty_inbox_is_valid_no_assets_fabricated(self):
            self.intake();self.assertEqual(m.read_json(self.out/'candidate_asset_registry.json')['candidates'],[])
            self.assertEqual(m.read_json(self.out/'candidate_intake_report.json')['status'],'EMPTY_INBOX_VALID')

    def test_batch_failure_does_not_partially_import(self):
            self.candidate();req=self.requests[3]['request_id'];(self.inbox/req/(req+'__A.png')).write_bytes(b'bad')
            with self.assertRaises(m.BridgeError):self.intake()
            self.assertFalse(self.out.exists())

    def test_packet_tampering_rejected_at_intake(self):
            p=self.package/'generation_packets/S001/generation_packet.json';d=m.read_json(p);d['prompt']+=' unsafe'
            p.write_text(json.dumps(d));self.candidate()
            with self.assertRaisesRegex(m.BridgeError,'PACKET_HASH_MISMATCH'):self.intake()

    def test_manifest_inbox_escape(self):
            self.manifest['authorized_inbox']=str(self.root)
            (self.package/'manual_handoff_manifest.json').write_text(json.dumps(self.manifest))
            with self.assertRaisesRegex(m.BridgeError,'CANDIDATE_PATH_ESCAPE'):self.intake()

    def test_output_outside_authorized_run_rejected(self):
            self.out=self.root/'bad_output'
            with self.assertRaisesRegex(m.BridgeError,'AUTHORIZED_EXTERNAL_OUTPUT_REQUIRED'):self.intake()

    def test_existing_output_not_overwritten(self):
            self.out.mkdir(parents=True)
            with self.assertRaisesRegex(m.BridgeError,'OUTPUT_ALREADY_EXISTS'):self.intake()

    def test_run_id_path_escape(self):
            with self.assertRaisesRegex(m.BridgeError,'INVALID_RUN_ID'):m.authorized_run('../outside',self.base)

    def test_floor_and_ambition_unchanged(self):
            before=copy.deepcopy(self.requests)
            for i in INITIAL:self.packet(i)
            self.assertEqual(before,self.requests)
            self.assertTrue(all(r['deck_quality_floor_status']=='VISUAL_QUALITY_BLOCKED' for r in self.requests.values()))

    def test_formal_source_unchanged_and_preview_labelled(self):
        before=m.sha_file(self.baseline);self.candidate();self.intake()
        self.assertEqual(before,m.sha_file(self.baseline))
        review=m.read_json(self.out/'review_manifest.json')
        self.assertFalse(review['approval_inferred'])
        self.assertEqual(review['preview_scientific_content_editability'],'RASTER_LOCAL_COMPARISON_ONLY')
        for name in ['candidate_comparison.pptx','candidate_comparison.pdf','candidate_contact_sheet.png']:
            self.assertTrue((self.out/'review'/name).is_file())

    def test_jpeg_webp_preview_png(self):
        self.candidate(suffix='.jpg');self.candidate(index=3,suffix='.webp');self.intake()
        for c in m.read_json(self.out/'candidate_asset_registry.json')['candidates']:
            with Image.open(self.out/'render_inputs'/(c['candidate_id']+'.png')) as image:self.assertEqual(image.format,'PNG')

if __name__=='__main__':unittest.main()
